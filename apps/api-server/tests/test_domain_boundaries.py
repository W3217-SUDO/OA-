"""领域路由迁移后，核对回款事务与用印文件补偿。"""

import asyncio
from datetime import date
import hashlib
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import IsolatedAsyncioTestCase
from unittest.mock import AsyncMock, patch
from types import SimpleNamespace
from copy import deepcopy

from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from fastapi import HTTPException, UploadFile

from app.areas.legal.router import import_business_records
from app.areas.legal import seals
from app.core.incoming_settlement import _active_settlements_by_receipt, _revert_incoming_allocation
from app.database import Base
from app.models import BusinessRecord, FileAttachment, FinanceTransaction, ReceivablePlan, User
from app.models_shared import AttachmentBatchInput
from app.areas.personal_agent import service as personal_service
from app.areas.personal_agent.confirmation_history import begin_personal_decision, project_personal_actions, read_confirmation_detail, read_confirmation_history
from app.areas.personal_agent.action_presentation import dismiss_personal_actions
from app.areas.personal_agent import conversations
from app.areas.personal_agent import identity_storage
from app.agent_mcp import auth as tool_auth
from app.agent_mcp import service as tool_service
from app.core import generated_file_storage as generated_storage
from app.core.generated_file_records import FILE_MODULE, SPACE_MODULE
from app.areas.personal_agent.ai_space_service import list_space_files, space_file_record, space_file_bytes
from sqlalchemy import func, select


class GeneratedFileStorageBoundaryTest(IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.db = self.sessions()
        self.users = [User(username=name, display_name="同名员工", role=role, role_ids=[role],
                           password_hash="unused", is_active=True)
                      for name, role in (("codex-ai-a", "user"), ("codex-ai-b", "user"), ("codex-ai-admin", "admin"))]
        self.db.add_all(self.users)
        await self.db.commit()
        self.actor = self.users[0]
        self.objects = {}
        self.patches = [patch.object(generated_storage, "_load_credentials", return_value=None),
                        patch("app.agent_mcp.auth.authenticated_identity", new=self.identity),
                        patch.object(generated_storage, "oss_put_editor_revision", new=AsyncMock(side_effect=self.put)),
                        patch.object(generated_storage, "read_generated_bytes", new=AsyncMock(side_effect=self.read)),
                        patch.object(generated_storage, "resolve_generated_scope", wraps=generated_storage.resolve_generated_scope),
                        patch.object(personal_service.settings, "ai_oss_bucket", "codex-ai-test"),
                        patch.object(personal_service.settings, "ai_oss_prefix", "CODEX/AI空间"),
                        patch.object(personal_service.settings, "ai_oss_endpoint", "https://oss.invalid"),
                        patch.object(personal_service.settings, "ai_oss_credentials_file", "isolated-ai-credentials.csv")]
        for item in self.patches:
            item.start()

    async def asyncTearDown(self):
        for item in reversed(self.patches):
            item.stop()
        await self.db.close()
        await self.engine.dispose()

    async def identity(self, db=None):
        return {"username": self.actor.username, "role": self.actor.role, "_actual_role_ids": self.actor.role_ids}

    async def put(self, bucket, key, content, content_type, *, endpoint):
        plan = await self.db.scalar(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE)
                                    .order_by(BusinessRecord.id.desc()))
        self.assertEqual(plan.status, "uploading")
        self.assertEqual(plan.data["path"], f"oss://{bucket}/{key}")
        self.assertEqual(plan.data["endpoint"], endpoint)
        self.objects[plan.data["path"]] = content
        return True

    async def read(self, item, *, max_bytes=20 * 1024 * 1024):
        return self.objects[item.path]

    async def save(self, **overrides):
        args = {"name": "真实结果.xls", "content": b"actual-xls-content", "content_type": "application/vnd.ms-excel",
                "source_kind": "mcp_export", "conversation_id": "source-conversation"}
        args.update(overrides)
        return await generated_storage.save_generated_file(self.db, await self.identity(), **args)

    async def office_edit(self, *, start=True):
        from docx import Document
        from app.areas.personal_agent import ai_space_service, dingtalk_office_service
        from app.core.constants import WORD_DOCUMENT_CONTENT_TYPE

        def word_bytes(text):
            document = Document()
            document.add_heading(text, level=1)
            document.add_table(rows=1, cols=1).cell(0, 0).text = "保留原生 Word 表格，不走文本 codec"
            output = BytesIO()
            document.save(output)
            return output.getvalue()

        self.actor.profile = {"dingtalk_user_id": "CODEX-office-member"}
        await self.db.commit()
        original, changed = word_bytes("原 Word"), word_bytes("修改后的 Word")
        result = await self.save(name="原生文稿.docx", content=original, content_type=WORD_DOCUMENT_CONTENT_TYPE)
        owner_id = self.actor.id
        app_digest = hashlib.sha256(b"CODEX-office-app").hexdigest()
        space = {"id": "CODEX-space", "ownerId": "CODEX-space-owner", "ownerType": "USER",
                 "creatorId": "CODEX-union", "corpId": "CODEX-corp", "scene": dingtalk_office_service.SCENE,
                 "sceneId": f"personal{owner_id}{app_digest[:16]}", "status": "NORMAL"}
        remote = SimpleNamespace(
            member_union_id=AsyncMock(return_value="CODEX-union"),
            storage_request=AsyncMock(return_value={"space": space}),
            upload_copy=AsyncMock(return_value={"id": "CODEX-dentry"}),
            dentry_info=AsyncMock(return_value={"version": "1", "size": len(original)}),
            download_version=AsyncMock(return_value=original),
            editor_url=AsyncMock(return_value="https://alidocs.dingtalk.com/CODEX-office"),
        )
        patches = [patch.object(personal_service.settings, "dingtalk_corp_id", "CODEX-corp"),
                   patch.object(personal_service.settings, "dingtalk_app_key", "CODEX-office-app"),
                   patch.object(ai_space_service, "read_generated_bytes", new=AsyncMock(side_effect=self.read))]
        patches.extend(patch.object(dingtalk_office_service.office, name, new=operation)
                       for name, operation in vars(remote).items())
        for item in patches:
            item.start()
            self.patches.append(item)
        session = None
        if start:
            session = await dingtalk_office_service.start_office_edit(
                result["file"]["id"], result["file"]["sha256"], await self.identity(), self.db)
        remote.dentry_info.return_value = {"version": "2", "size": len(changed)}
        remote.download_version.return_value = changed
        return result, session, remote, changed

    async def test_real_names_stable_owners_and_one_attachment(self):
        first = await self.save()
        self.assertEqual(first["file"]["status"], "available")
        self.assertEqual(first["file"]["owner_user_id"], self.actor.id)
        self.assertIn(f"/同名员工--u{self.actor.id}/", first["attachment"].path)
        self.assertTrue(first["attachment"].path.endswith("/真实结果.xls"))
        self.assertEqual(first["file"]["attachment_id"], first["attachment"].id)
        self.assertEqual(first["file"]["content_type"], "application/vnd.ms-excel")
        old_folder = first["attachment"].path.split("/" + first["file"]["id"] + "/")[0]
        self.actor.display_name = "改名员工"
        await self.db.commit()
        second = await self.save(name="第二结果.xls")
        self.assertTrue(second["attachment"].path.startswith(old_folder + "/"))
        self.actor = self.users[1]
        third = await self.save()
        self.assertIn(f"/同名员工--u{self.actor.id}/", third["attachment"].path)
        self.assertNotEqual(first["attachment"].path, third["attachment"].path)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 3)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(BusinessRecord)
                                             .where(BusinessRecord.module == SPACE_MODULE)), 2)

    async def test_zero_byte_case_text_create_edit_and_generic_upload_are_preserved(self):
        from app.api_schemas.case_documents import CaseAiDraftCreateInput, CaseAiDraftUpdateInput
        from app.areas.legal.case_documents import create_case_ai_draft, update_case_ai_draft_content
        from app.areas.system.attachments import upload_attachment
        from app.core.constants import AI_SPACE_CATEGORY

        self.actor = self.users[2]
        case = BusinessRecord(module="case", serial_no="CODEX-AI-ZERO", title="零字节兼容案件",
                              owner=self.actor.username, status="草稿", data={})
        self.db.add(case)
        await self.db.commit()
        empty_sha = hashlib.sha256(b"").hexdigest()
        for suffix in ("txt", "md"):
            created = await create_case_ai_draft(case.id, CaseAiDraftCreateInput(name=f"空白.{suffix}", content=""),
                                                 identity=await self.identity(), db=self.db)
            self.assertEqual(created["size"], 0)
            item = await self.db.get(FileAttachment, created["id"])
            old_path = item.path
            edited = await update_case_ai_draft_content(case.id, item.id, CaseAiDraftUpdateInput(content=""),
                                                       identity=await self.identity(), db=self.db)
            self.assertEqual(edited["size"], 0)
            self.assertEqual(self.objects[old_path], b"")
        with patch("app.core.case_automation.advance_case_for_uploaded_document", new=AsyncMock()), \
                patch("app.core.documents._sync_case_document_readiness", new=AsyncMock()):
            uploaded = await upload_attachment(
                file=UploadFile(filename="原空文件.txt", file=BytesIO(b"")), record_id=case.id,
                finance_transaction_id=None, source_case_id=None, customer_guid=None,
                is_license=None, document_date=None, category=AI_SPACE_CATEGORY, remark="",
                identity=await self.identity(), db=self.db,
            )
        self.assertEqual(uploaded["size"], 0)
        records = (await self.db.scalars(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE))).all()
        self.assertEqual(len(records), 5)
        for record in records:
            self.assertEqual(record.status, "available")
            self.assertEqual(record.data["size"], 0)
            self.assertEqual(record.data["sha256"], empty_sha)
            self.assertEqual(self.objects[record.data["path"]], b"")

    async def test_bound_employee_is_authoritative_and_conflicting_binding_rejected(self):
        employee = BusinessRecord(module="hr", serial_no="CODEX-AI-HR", title="真实档案姓名", owner=self.actor.username,
                                  data={"username": self.actor.username, "system_user_id": self.actor.id})
        self.db.add(employee)
        await self.db.commit()
        result = await self.save()
        self.assertEqual(result["file"]["owner_name"], "真实档案姓名")
        self.assertIn("/真实档案姓名--u", result["attachment"].path)
        employee.data = {"username": self.actor.username, "system_user_id": self.users[1].id}
        await self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            await self.save()
        self.assertEqual(raised.exception.status_code, 409)
        self.assertEqual(len(self.objects), 1)

    async def test_missing_config_or_name_never_uploads_or_falls_back(self):
        with patch.object(personal_service.settings, "ai_oss_bucket", ""):
            with self.assertRaises(HTTPException) as raised:
                await self.save()
            self.assertEqual(raised.exception.status_code, 503)
        self.actor.display_name = ""
        await self.db.commit()
        with self.assertRaises(HTTPException) as raised:
            await self.save()
        self.assertEqual(raised.exception.status_code, 422)
        self.assertFalse(self.objects)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)

    async def test_failed_readback_is_tracked_without_available_attachment(self):
        with patch.object(generated_storage, "read_generated_bytes", new=AsyncMock(return_value=b"changed")):
            with self.assertRaises(HTTPException) as raised:
                await self.save()
        self.assertEqual(raised.exception.status_code, 409)
        plan = await self.db.scalar(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE))
        self.assertEqual(plan.status, "uncertain")
        self.assertIn(plan.data["path"], self.objects)
        self.assertIsNone(plan.data["attachment_id"])
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)

    async def test_callback_failure_keeps_plan_and_rolls_back_final_attachment(self):
        callback = AsyncMock(side_effect=RuntimeError("isolated final transaction failure"))
        with self.assertRaises(HTTPException):
            await self.save(on_available=callback)
        plan = await self.db.scalar(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE))
        self.assertEqual(plan.status, "uncertain")
        self.assertEqual(plan.data["error"], "文件保存未完成（RuntimeError），请人工核对")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)
        self.assertEqual(len(self.objects), 1)

    async def test_owner_access_rejects_page_admin_but_real_admin_reads(self):
        result = await self.save()
        owner_id = self.actor.id
        self.actor = self.users[1]
        spoof = {"username": self.actor.username, "role": "admin", "_actual_role_ids": ["admin"], "_page_menu_capability": True}
        with self.assertRaises(HTTPException) as raised:
            await space_file_record(result["file"]["id"], spoof, self.db)
        self.assertEqual(raised.exception.status_code, 403)
        self.assertEqual((await list_space_files(await self.identity(), self.db, page=1, page_size=20))["total"], 0)
        self.actor = self.users[2]
        item = await space_file_record(result["file"]["id"], await self.identity(), self.db)
        self.assertEqual(item.data["owner_user_id"], owner_id)
        listing = await list_space_files(await self.identity(), self.db, page=1, page_size=20, owner_user_id=owner_id)
        self.assertEqual(listing["total"], 1)
        self.assertFalse(listing["items"][0]["can_delete"], "管理员只读他人目录不投影代删能力")
        from app.core.permissions import _require_record_module_menu
        from app.core.system import _record_module_menu_allowed
        with self.assertRaises(HTTPException) as raised:
            await _require_record_module_menu(FILE_MODULE, await self.identity(), self.db, action="编辑")
        self.assertEqual(raised.exception.status_code, 403)
        self.assertFalse(_record_module_menu_allowed(FILE_MODULE, spoof, {}))

    async def test_download_checks_real_bytes_and_name_validation(self):
        for name in ("../x.xls", "a\\b.xls", "x\x00.xls", "C:x.xls"):
            with self.assertRaises(HTTPException):
                generated_storage.safe_original_name(name)
        result = await self.save()
        with patch("app.areas.personal_agent.ai_space_service.read_generated_bytes", new=AsyncMock(return_value=b"actual-xls-content")):
            item, content = await space_file_bytes(result["file"]["id"], await self.identity(), self.db)
        self.assertEqual(item["id"], result["file"]["id"])
        self.assertEqual(content, b"actual-xls-content")
        with patch("app.areas.personal_agent.ai_space_service.read_generated_bytes", new=AsyncMock(return_value=b"wrong")):
            with self.assertRaises(HTTPException) as raised:
                await space_file_bytes(result["file"]["id"], await self.identity(), self.db)
        self.assertEqual(raised.exception.status_code, 409)

    async def test_case_attachment_is_indexed_once_and_revisions_keep_original_bytes(self):
        self.actor = self.users[2]
        case = BusinessRecord(module="case", serial_no="CODEX-AI-CASE", title="真实案件名称", owner=self.actor.username,
                              status="草稿", data={})
        self.db.add(case)
        await self.db.commit()
        with patch("app.core.permissions._require_case_attachment_upload_access", new=AsyncMock()):
            first = await self.save(case_record=case, source_kind="case_ai_draft")
            old_path = first["attachment"].path
            second = await self.save(case_record=case, attachment=first["attachment"], content=b"second-revision",
                                     source_kind="case_ai_draft")
        self.assertEqual(first["attachment"].id, second["attachment"].id)
        self.assertEqual(second["attachment"].record_id, case.id)
        self.assertEqual(second["file"]["case_id"], case.id)
        self.assertIn(f"/CODEX-AI-CASE-真实案件名称--c{case.id}/", second["attachment"].path)
        self.assertNotEqual(old_path, second["attachment"].path)
        self.assertEqual(self.objects[old_path], b"actual-xls-content")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 1)
        listing = await list_space_files(await self.identity(), self.db, page=1, page_size=20)
        self.assertEqual(listing["total"], 2)
        self.assertEqual({item["attachment_id"] for item in listing["items"]}, {second["attachment"].id})

    async def test_personal_edit_links_new_id_reads_changed_content_and_preserves_original_on_failure(self):
        from app.areas.personal_agent import ai_space_service

        original = await self.save(name="原文件.txt", content="原始正文".encode(), content_type="text/plain")
        original_path = original["attachment"].path
        with patch.object(ai_space_service, "read_generated_bytes", new=AsyncMock(side_effect=self.read)):
            edited = await ai_space_service.edit_space_file(
                original["file"]["id"], "修改版本.txt", "修改后的正文", await self.identity(), self.db)
            self.assertNotEqual(edited["file"]["id"], original["file"]["id"])
            self.assertEqual(edited["supersedes"], original["file"]["id"])
            self.assertEqual(edited["file"]["supersedes"], original["file"]["id"])
            self.assertTrue(edited["file"]["can_delete"])
            new_text = await ai_space_service.space_file_text(edited["file"]["id"], await self.identity(), self.db)
            old_text = await ai_space_service.space_file_text(original["file"]["id"], await self.identity(), self.db)
            self.assertEqual(new_text["id"], edited["file"]["id"])
            self.assertEqual(new_text["content"], "修改后的正文")
            self.assertEqual(old_text["content"], "原始正文")
            version = await space_file_record(edited["file"]["id"], await self.identity(), self.db)
            self.assertEqual(version.data["supersedes"], original["file"]["id"])
            with patch.object(generated_storage, "save_generated_file", new=AsyncMock(side_effect=HTTPException(503, "保存未确认"))) as save:
                with self.assertRaises(HTTPException) as raised:
                    await ai_space_service.edit_space_file(
                        original["file"]["id"], "失败版本.txt", "不能覆盖原文", await self.identity(), self.db)
                self.assertEqual(raised.exception.status_code, 503)
                save.assert_awaited_once()
            for actor in self.users[1:]:
                self.actor = actor
                with self.assertRaises(HTTPException) as raised:
                    await ai_space_service.edit_space_file(
                        original["file"]["id"], "越权.txt", "不允许写入", await self.identity(), self.db)
                self.assertEqual(raised.exception.status_code, 403, "真实管理员也不能代编辑他人版本")
        self.assertEqual(self.objects[original_path], "原始正文".encode())
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 2)

    async def test_personal_delete_is_exact_idempotent_and_preserves_other_versions(self):
        from app.areas.personal_agent import ai_space_deletion, ai_space_service

        original = await self.save(name="旧版本.txt", content=b"original", content_type="text/plain")
        with patch.object(ai_space_service, "read_generated_bytes", new=AsyncMock(side_effect=self.read)):
            edited = await ai_space_service.edit_space_file(
                original["file"]["id"], "新版本.txt", "edited", await self.identity(), self.db)
        target = edited["file"]
        target_record = await space_file_record(target["id"], await self.identity(), self.db)
        target_path = target_record.data["path"]

        async def remove_exact(plan):
            record = await self.db.get(BusinessRecord, target_record.id, populate_existing=True)
            self.assertEqual(record.status, "deleting")
            self.assertEqual(record.data["delete_plan"]["operation_id"], plan["operation_id"])
            self.assertEqual(plan["path"], target_path)
            self.assertEqual(plan["sha256"], hashlib.sha256(self.objects[target_path]).hexdigest())
            self.assertEqual(plan["size"], len(self.objects[target_path]))
            self.objects.pop(target_path)
            return {"blob_status": "absent", "historical_versions_purged": False}

        result = await ai_space_deletion.delete_space_file(target["id"], target["sha256"], await self.identity(), self.db)
        again = await ai_space_deletion.delete_space_file(target["id"], target["sha256"], await self.identity(), self.db)
        self.assertEqual(result, {"id": target["id"], "status": "deleted", "blob_status": "retained"})
        self.assertEqual(again, result)
        tombstone = await self.db.get(BusinessRecord, target_record.id, populate_existing=True)
        self.assertEqual(tombstone.status, "deleted")
        self.assertEqual(tombstone.data["supersedes"], original["file"]["id"])
        self.assertEqual(tombstone.data["delete_plan"]["blob_status"], "retained")
        self.assertIn(target_path, self.objects)
        self.assertEqual(self.objects[original["attachment"].path], b"original")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 2)
        listing = await list_space_files(await self.identity(), self.db, page=1, page_size=20)
        self.assertEqual([item["id"] for item in listing["items"]], [original["file"]["id"]])
        for read in (space_file_record, space_file_bytes):
            with self.assertRaises(HTTPException) as raised:
                await read(target["id"], await self.identity(), self.db)
            self.assertEqual(raised.exception.status_code, 410)

    async def test_personal_delete_rejects_other_owners_case_identity_and_changed_hash_before_blob(self):
        from app.areas.personal_agent import ai_space_deletion

        owner = self.actor
        original = await self.save()
        with patch.object(ai_space_deletion, "delete_personal_blob", new=AsyncMock()) as delete:
            for actor in self.users[1:]:
                self.actor = actor
                with self.assertRaises(HTTPException) as raised:
                    await ai_space_deletion.delete_space_file(
                        original["file"]["id"], original["file"]["sha256"], await self.identity(), self.db)
                self.assertEqual(raised.exception.status_code, 403)
            self.actor = owner
            with self.assertRaises(HTTPException) as raised:
                await ai_space_deletion.delete_space_file(original["file"]["id"], "0" * 64, await self.identity(), self.db)
            self.assertEqual(raised.exception.status_code, 409)
            identity_file = await self.save(name="IDENTITY.md", source_kind="personal_identity")
            with self.assertRaises(HTTPException) as raised:
                await ai_space_deletion.delete_space_file(
                    identity_file["file"]["id"], identity_file["file"]["sha256"], await self.identity(), self.db)
            self.assertEqual(raised.exception.status_code, 409)
            case = BusinessRecord(module="case", serial_no="CODEX-AI-NO-DELETE", title="隔离案件", owner=owner.username, status="草稿", data={})
            self.db.add(case)
            await self.db.commit()
            with patch("app.core.permissions._require_case_attachment_upload_access", new=AsyncMock()):
                case_file = await self.save(case_record=case, source_kind="case_draft")
            with self.assertRaises(HTTPException) as raised:
                await ai_space_deletion.delete_space_file(
                    case_file["file"]["id"], case_file["file"]["sha256"], await self.identity(), self.db)
            self.assertEqual(raised.exception.status_code, 409)
            delete.assert_not_awaited()
        self.assertEqual(len(self.objects), 3)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 3)

    async def test_personal_delete_is_logical_and_retains_attachment_for_cleanup(self):
        from app.areas.personal_agent import ai_space_deletion

        original = await self.save()
        file = original["file"]
        original_path = original["attachment"].path
        identity = await self.identity()
        result = await ai_space_deletion.delete_space_file(file["id"], file["sha256"], identity, self.db)
        self.assertEqual(result, {"id": file["id"], "status": "deleted", "blob_status": "retained"})
        record = await self.db.scalar(select(BusinessRecord).where(
            BusinessRecord.module == FILE_MODULE, BusinessRecord.serial_no == "AIGF-" + file["id"],
        ).execution_options(populate_existing=True))
        self.assertEqual(record.status, "deleted")
        self.assertEqual(record.data["delete_plan"]["blob_status"], "retained")
        self.assertEqual(self.objects[original_path], b"actual-xls-content")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 1)
        listing = await list_space_files(identity, self.db, page=1, page_size=20)
        self.assertEqual(listing["total"], 0)

    async def test_office_new_and_unchanged_versions_project_owner_capability_and_link_in_final_transaction(self):
        from app.areas.personal_agent import dingtalk_office_service as office

        source, session, remote, changed = await self.office_edit()
        source_record_id = source["attachment"].record_id
        original_path = source["attachment"].path
        original_bytes = self.objects[original_path]
        commit = self.db.commit
        linked_before_commit = []

        async def inspect_final_commit():
            versions = (await self.db.scalars(select(BusinessRecord).where(
                BusinessRecord.module == FILE_MODULE,
                BusinessRecord.data["source_kind"].as_string() == "dingtalk_office_edit",
            ))).all()
            for version in versions:
                if version.status == "available":
                    parent = await self.db.get(BusinessRecord, source_record_id)
                    self.assertEqual(version.data["supersedes"], source["file"]["id"])
                    self.assertEqual(parent.data[office.SESSION_KEY]["result_file_id"], version.data["file_id"])
                    linked_before_commit.append(version.data["file_id"])
            await commit()

        with patch.object(self.db, "commit", new=inspect_final_commit):
            saved = await office.save_office_edit(source["file"]["id"], session["session_id"], await self.identity(), self.db)
        self.assertFalse(saved["unchanged"])
        self.assertNotEqual(saved["file"]["id"], source["file"]["id"])
        self.assertEqual(saved["file"]["supersedes"], source["file"]["id"])
        self.assertEqual(saved["supersedes"], source["file"]["id"])
        self.assertTrue(saved["file"]["can_delete"])
        self.assertEqual(linked_before_commit, [saved["file"]["id"]])
        self.assertEqual(saved["file"]["owner_user_id"], self.actor.id)
        self.assertIsNone(saved["file"]["case_id"])
        self.assertEqual(saved["file"]["sha256"], hashlib.sha256(changed).hexdigest())
        _, downloaded = await space_file_bytes(saved["file"]["id"], await self.identity(), self.db)
        self.assertEqual(downloaded, changed)
        again = await office.save_office_edit(source["file"]["id"], session["session_id"], await self.identity(), self.db)
        self.assertTrue(again["unchanged"])
        self.assertEqual(again["file"]["id"], saved["file"]["id"])
        self.assertTrue(again["file"]["can_delete"])
        self.assertEqual(remote.download_version.await_count, 2, "同一钉钉版本不重复下载或生成 OA 文件")
        from app.data_models.records import WorkflowEvent
        events = (await self.db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == source_record_id))).all()
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0].action, "钉钉在线编辑自动回写")
        self.assertIn(saved["file"]["id"], events[0].comment)
        self.assertEqual(self.objects[original_path], original_bytes)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 2)

    async def test_office_source_and_saved_result_tombstones_block_without_replay(self):
        from app.areas.personal_agent import dingtalk_office_service as office

        source, session, remote, _ = await self.office_edit()
        identity = await self.identity()
        saved = await office.save_office_edit(source["file"]["id"], session["session_id"], identity, self.db)
        records = {item: await self.db.scalar(select(BusinessRecord).where(
            BusinessRecord.module == FILE_MODULE, BusinessRecord.serial_no == "AIGF-" + item,
        )) for item in (source["file"]["id"], saved["file"]["id"])}
        provider_uploads = remote.upload_copy.await_count
        object_count = len(self.objects)
        for target, record in records.items():
            available_data = deepcopy(record.data)
            for status in ("deleted", "deleting", "delete_failed", "delete_uncertain"):
                with self.subTest(target=target, status=status):
                    plan = {"operation_id": "CODEX-delete-" + target, "blob_status": "absent" if status == "deleted" else "pending"}
                    record.status = status
                    record.data = {**record.data, "delete_plan": plan}
                    await self.db.commit()
                    if target == source["file"]["id"]:
                        with self.assertRaises(HTTPException) as raised:
                            await office.start_office_edit(target, source["file"]["sha256"], identity, self.db)
                        self.assertEqual(raised.exception.status_code, 410 if status == "deleted" else 409)
                    with self.assertRaises(HTTPException) as raised:
                        await office.save_office_edit(source["file"]["id"], session["session_id"], identity, self.db)
                    self.assertEqual(raised.exception.status_code, 410 if status == "deleted" else 409)
                    self.assertEqual(record.data["delete_plan"], plan)
                    self.assertEqual(remote.upload_copy.await_count, provider_uploads)
                    self.assertEqual(len(self.objects), object_count, "墓碑结果不能切回源文件或另造 OA 新版本")
            record.status = "available"
            record.data = available_data
            await self.db.commit()

    async def test_office_owner_case_binding_and_expected_hash_reject_before_copy(self):
        from app.areas.personal_agent import dingtalk_office_service as office

        source, session, remote, _ = await self.office_edit()
        owner = self.actor
        for actor in self.users[1:]:
            self.actor = actor
            for operation in ("open", "save"):
                with self.subTest(actor=actor.username, operation=operation), self.assertRaises(HTTPException) as raised:
                    if operation == "open":
                        await office.start_office_edit(source["file"]["id"], source["file"]["sha256"], await self.identity(), self.db)
                    else:
                        await office.save_office_edit(source["file"]["id"], session["session_id"], await self.identity(), self.db)
                self.assertEqual(raised.exception.status_code, 403)
        self.actor = owner
        with self.assertRaises(HTTPException) as raised:
            await office.start_office_edit(source["file"]["id"], "0" * 64, await self.identity(), self.db)
        self.assertEqual(raised.exception.status_code, 409)
        self.actor = self.users[2]
        case = BusinessRecord(module="case", serial_no="CODEX-OFFICE-CASE", title="可见案件", owner=self.actor.username, status="草稿", data={})
        self.db.add(case)
        await self.db.commit()
        with patch("app.core.permissions._require_case_attachment_upload_access", new=AsyncMock()):
            case_file = await self.save(case_record=case, source_kind="case_draft")
        with self.assertRaises(HTTPException) as raised:
            await office.start_office_edit(case_file["file"]["id"], case_file["file"]["sha256"], await self.identity(), self.db)
        self.assertEqual(raised.exception.status_code, 403)
        self.assertEqual(remote.upload_copy.await_count, 1)

    async def test_office_copy_failure_preserves_http_unknown_and_cancellation_with_source_tombstone(self):
        from app.areas.personal_agent import dingtalk_office_service as office

        for failure in (HTTPException(403, "原权限拒绝"), HTTPException(409, "原版本拒绝"),
                        HTTPException(410, "原删除拒绝"), RuntimeError("原未知错误"), asyncio.CancelledError()):
            with self.subTest(failure=type(failure).__name__, status=getattr(failure, "status_code", None)):
                source, _, remote, _ = await self.office_edit(start=False)
                source_record_id = source["attachment"].record_id
                plan = {"operation_id": "CODEX-copy-delete", "blob_status": "absent"}

                async def delete_during_upload(*args):
                    record = await self.db.get(BusinessRecord, source_record_id, populate_existing=True)
                    record.status = "deleted"
                    record.data = {**record.data, "delete_plan": plan}
                    await self.db.commit()
                    raise failure

                remote.upload_copy.side_effect = delete_during_upload
                with self.assertRaises(type(failure)) as raised:
                    await office.start_office_edit(source["file"]["id"], source["file"]["sha256"], await self.identity(), self.db)
                self.assertIs(raised.exception, failure, "墓碑审计不能把原异常替换成普通 available 锁 409")
                await self.db.refresh(self.actor)
                record = await self.db.get(BusinessRecord, source_record_id, populate_existing=True)
                self.assertEqual(record.status, "deleted")
                self.assertEqual(record.data["delete_plan"], plan)
                self.assertEqual(record.data["sha256"], source["file"]["sha256"])
                self.assertEqual(record.data["source_kind"], source["file"]["source_kind"])

    async def test_office_writer_uncertain_handle_survives_deleted_parent_and_audit_failure(self):
        from app.areas.personal_agent import dingtalk_office_service as office

        for fail_audit in (False, True):
            with self.subTest(fail_audit=fail_audit):
                source, session, _, changed = await self.office_edit()
                source_record_id = source["attachment"].record_id
                plan = {"operation_id": "CODEX-writer-delete", "blob_status": "absent"}

                async def delete_before_on_available(*args, **kwargs):
                    await self.put(*args, **kwargs)
                    record = await self.db.get(BusinessRecord, source_record_id, populate_existing=True)
                    record.status = "deleted"
                    record.data = {**record.data, "delete_plan": plan}
                    await self.db.commit()
                    return True

                original_audit = office._uncertain

                async def audit_with_database_failure(*args):
                    # 只注入审计读取失败，保留真实审计函数的异常记录与原错误保真。
                    with patch.object(self.db, "scalar", new=AsyncMock(side_effect=RuntimeError("隔离审计失败"))):
                        await original_audit(*args)

                with patch.object(generated_storage, "oss_put_editor_revision", new=AsyncMock(side_effect=delete_before_on_available)), \
                        patch.object(office.logger, "exception") as audit_log, \
                        patch.object(office, "_uncertain", new=audit_with_database_failure if fail_audit else original_audit):
                    with self.assertRaises(HTTPException) as raised:
                        await office.save_office_edit(source["file"]["id"], session["session_id"], await self.identity(), self.db)
                if fail_audit:
                    audit_log.assert_called_once()
                    self.assertTrue(any("RuntimeError" in note for note in raised.exception.__notes__))
                else:
                    audit_log.assert_not_called()
                self.assertEqual(raised.exception.status_code, 409)
                self.assertEqual(raised.exception.detail["status"], "uncertain")
                file_id = raised.exception.detail["file_id"]
                self.assertNotEqual(file_id, source["file"]["id"])
                await self.db.refresh(self.actor)
                version = await self.db.scalar(select(BusinessRecord).where(
                    BusinessRecord.module == FILE_MODULE, BusinessRecord.serial_no == "AIGF-" + file_id,
                ).execution_options(populate_existing=True))
                self.assertEqual(version.status, "uncertain")
                self.assertIsNone(version.data["attachment_id"])
                self.assertEqual(version.data["sha256"], hashlib.sha256(changed).hexdigest())
                record = await self.db.get(BusinessRecord, source_record_id, populate_existing=True)
                self.assertEqual(record.status, "deleted")
                self.assertEqual(record.data["delete_plan"], plan)

    async def test_final_commit_failure_keeps_remote_object_and_uncertain_plan(self):
        commit = self.db.commit
        calls = 0

        async def fail_final_commit():
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("isolated final commit failed")
            await commit()

        with patch.object(self.db, "commit", new=fail_final_commit):
            with self.assertRaises(HTTPException):
                await self.save()
        self.assertEqual(calls, 3)
        plan = await self.db.scalar(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE))
        self.assertEqual(plan.status, "uncertain")
        self.assertIn(plan.data["path"], self.objects)
        self.assertIsNone(plan.data["attachment_id"])
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)

    async def test_cancelled_upload_keeps_uncertain_plan(self):
        import asyncio

        with patch.object(generated_storage, "oss_put_editor_revision", new=AsyncMock(side_effect=asyncio.CancelledError())):
            with self.assertRaises(asyncio.CancelledError):
                await self.save()
        plan = await self.db.scalar(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE))
        self.assertEqual(plan.status, "uncertain")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)

    async def test_no_case_message_capability_and_true_message_save(self):
        from app.areas.personal_agent.ai_space_service import save_space_message

        identity = {**await self.identity(), "display_name": self.actor.display_name}
        state = {"conversation_id": "original", "messages": [], "pending_actions": []}
        with patch.object(personal_service, "PersonalMcpToolSession", wraps=personal_service.PersonalMcpToolSession) as tool_session, \
                patch("app.core.oss_attachments._load_credentials", return_value=None), \
                patch.object(personal_service, "_agent_skill_for_identity", new=AsyncMock(return_value=SimpleNamespace(id="test", name="测试技能"))), \
                patch.object(personal_service, "_agent_message_skill_for_identity", new=AsyncMock(return_value=(SimpleNamespace(id="test", name="测试技能"), "生成文书"))), \
                patch.object(personal_service, "prepare_materials", new=AsyncMock(return_value={"case": None, "readings": [], "can_save_document": False})), \
                patch.object(personal_service, "build_context", new=AsyncMock(return_value={})), \
                patch.object(personal_service, "_prompt_messages", return_value=[]), \
                patch.object(personal_service, "request_model", new=AsyncMock(return_value="真实完整回答")) as model_request, \
                patch.object(personal_service, "_save_state", new=AsyncMock()):
            result = await personal_service._generate_response(identity, self.db, "生成文书", AsyncMock(), state=state,
                                                               skill_id="test", attachment_ids=[], case_id=None, document_ids=[],
                                                               response_mode="fast", on_progress=None)
        message = result["state"]["messages"][-1]
        tool_session.assert_called_once_with(identity, self.db, conversation_id="original")
        model_request.assert_awaited_once()
        tools = model_request.await_args.kwargs["tools"]
        self.assertIsInstance(tools, personal_service.PersonalMcpToolSession)
        self.assertIs(model_request.await_args.kwargs["model_tools"], tools.model_tools)
        self.assertTrue(message["can_save_document"])
        self.assertTrue(message["can_save_personal_document"])
        self.assertFalse(message["can_save_case_document"])
        with TemporaryDirectory(prefix="oa-test-personal-message-") as temporary:
            directory = Path(temporary)
            snapshot = await conversations.create_conversation(directory)
            snapshot["state"]["messages"] = [message]
            await conversations.save_state(directory, snapshot["state"])
            body = SimpleNamespace(conversation_id=snapshot["state"]["conversation_id"], message_index=0,
                                   name="真实文书.docx", expected_content_sha256=hashlib.sha256(message["content"].encode()).hexdigest())
            with patch.object(personal_service, "_user_dir", return_value=directory):
                saved = await save_space_message(body, identity, self.db)
                self.assertEqual(saved["source_kind"], "assistant_message")
                self.assertEqual(saved["conversation_id"], body.conversation_id)
                body.expected_content_sha256 = "0" * 64
                with self.assertRaises(HTTPException) as raised:
                    await save_space_message(body, identity, self.db)
                self.assertEqual(raised.exception.status_code, 409)

    async def test_ai_endpoint_is_separate_for_signed_download_and_missing_config(self):
        from urllib.parse import urlsplit
        from app.core.oss_attachments import oss_attachment_signed_url

        old = SimpleNamespace(path="oss://historical-test/old.docx")
        new = SimpleNamespace(path="oss://codex-ai-test/CODEX/AI空间/new.docx")
        with patch("app.core.oss_attachments._load_credentials", return_value=("isolated-id", "isolated-secret")) as load, \
                patch.object(personal_service.settings, "oss_endpoint", "https://oss-old.invalid"), \
                patch.object(personal_service.settings, "ai_oss_endpoint", "https://oss-ai.invalid"), \
                patch.object(personal_service.settings, "ai_oss_credentials_file", "explicit-ai-credentials.csv"):
            self.assertEqual(urlsplit(oss_attachment_signed_url(old)).hostname, "historical-test.oss-old.invalid")
            load.assert_called_with()
            self.assertEqual(urlsplit(oss_attachment_signed_url(new)).hostname, "codex-ai-test.oss-ai.invalid")
            load.assert_called_with("explicit-ai-credentials.csv")
            with patch.object(personal_service.settings, "ai_oss_endpoint", ""):
                with self.assertRaises(HTTPException) as raised:
                    oss_attachment_signed_url(new)
                self.assertEqual(raised.exception.status_code, 503)
                self.assertEqual(urlsplit(oss_attachment_signed_url(old)).hostname, "historical-test.oss-old.invalid")
        with patch.object(personal_service.settings, "ai_oss_credentials_file", ""), \
                patch.object(personal_service.settings, "oss_credentials_file", "must-not-inherit.csv"):
            with self.assertRaises(HTTPException) as raised:
                oss_attachment_signed_url(new)
            self.assertEqual(raised.exception.status_code, 503)

    async def test_unlinked_hr_owner_is_creator_not_employee_binding(self):
        self.db.add(BusinessRecord(module="hr", serial_no="CODEX-AI-UNLINKED", title="其他档案姓名",
                                   owner=self.actor.username, data={}))
        await self.db.commit()
        saved = await self.save()
        self.assertEqual(saved["file"]["owner_name"], self.actor.display_name)
        self.assertIn(f"/同名员工--u{self.actor.id}/", saved["attachment"].path)

    async def test_generic_case_ai_upload_preserves_shared_storage_and_workflow(self):
        from app.areas.system.attachments import upload_attachment
        from app.core.constants import AI_SPACE_CATEGORY

        self.actor = self.users[2]
        case = BusinessRecord(module="case", serial_no="CODEX-AI-GENERIC", title="隔离案件",
                              owner=self.actor.username, status="草稿", data={})
        self.db.add(case)
        await self.db.commit()
        file = UploadFile(filename="草稿.docx", file=BytesIO(b"uploaded-unchanged"))
        with patch("app.core.permissions._require_case_attachment_upload_access", new=AsyncMock()), \
                patch("app.core.case_automation.advance_case_for_uploaded_document", new=AsyncMock()) as advance, \
                patch("app.core.documents._sync_case_document_readiness", new=AsyncMock()) as readiness, \
                patch("app.areas.system.attachments.Path.write_bytes", side_effect=AssertionError("AI must not write local")):
            result = await upload_attachment(file=file, record_id=case.id, finance_transaction_id=None,
                                             source_case_id=None, customer_guid=None, is_license=None, document_date=None,
                                             category=AI_SPACE_CATEGORY, remark="原上传备注", identity=await self.identity(), db=self.db)
        self.assertEqual(result["storage_backend"], "oss")
        self.assertEqual(result["category"], AI_SPACE_CATEGORY)
        self.assertEqual(result["remark"], "原上传备注")
        advance.assert_awaited_once()
        readiness.assert_awaited_once()
        from app.data_models.records import WorkflowEvent
        event = await self.db.scalar(select(WorkflowEvent).where(WorkflowEvent.record_id == case.id))
        self.assertEqual(event.action, "上传归档材料")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 1)
        self.assertEqual(len(self.objects), 1)
        with patch.object(file, "read", new=AsyncMock(return_value=b"x" * (20 * 1024 * 1024 + 1))), \
                patch("app.core.permissions._require_case_attachment_upload_access", new=AsyncMock()), \
                patch("app.core.cases._save_case_ai_upload", new=AsyncMock()) as save:
            with self.assertRaises(HTTPException) as raised:
                await upload_attachment(file=file, record_id=case.id, finance_transaction_id=None,
                                        source_case_id=None, customer_guid=None, is_license=None, document_date=None,
                                        category=AI_SPACE_CATEGORY, remark="", identity=await self.identity(), db=self.db)
            self.assertEqual(raised.exception.status_code, 413)
            save.assert_not_awaited()

    async def test_configuration_ready_is_not_put_permission_proof_and_reads_do_not_create_space(self):
        from app.areas.personal_agent.ai_space_service import space_info

        with patch("app.core.oss_attachments._load_credentials", return_value=None):
            info = await space_info(await self.identity(), self.db)
            self.assertTrue(info["configured"])
            self.assertFalse(info["put_verified"])
            self.assertEqual(info["verification_scope"], "configuration_only")
            with patch.object(personal_service.settings, "ai_oss_endpoint", ""):
                with self.assertRaises(HTTPException) as raised:
                    await space_info(await self.identity(), self.db)
                self.assertEqual(raised.exception.status_code, 503)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(BusinessRecord)
                                             .where(BusinessRecord.module == SPACE_MODULE)), 0)
        self.assertFalse(self.objects)

    async def test_committed_available_state_is_not_rewritten_as_failed_when_ack_is_lost(self):
        commit = self.db.commit
        calls = 0

        async def commit_with_lost_ack():
            nonlocal calls
            calls += 1
            await commit()
            if calls == 2:
                raise RuntimeError("isolated committed response lost")

        with patch.object(self.db, "commit", new=commit_with_lost_ack):
            with self.assertRaises(HTTPException) as raised:
                await self.save()
        self.assertEqual(raised.exception.detail["status"], "available")
        self.assertIn("文件已持久提交为 available", raised.exception.detail["detail"])
        plan = await self.db.scalar(select(BusinessRecord).where(BusinessRecord.module == FILE_MODULE))
        self.assertEqual(plan.status, "available")
        self.assertIsNone(plan.data["error"])
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 1)


class PersonalIdentityStorageBoundaryTest(IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temporary = TemporaryDirectory(prefix="CODEX-personal-identity-")
        root = Path(self.temporary.name)
        self.engine = create_async_engine("sqlite+aiosqlite:///" + (root / "identity.sqlite").as_posix())
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.db = self.sessions()
        self.actor = User(username="CODEX-identity-owner", display_name="真实员工", role="user", role_ids=["user"],
                          password_hash="unused", is_active=True)
        self.db.add(self.actor)
        await self.db.commit()
        self.username = self.actor.username
        self.objects = {}
        self.writer = AsyncMock(side_effect=self.put)
        self.patches = [patch.object(generated_storage, "_load_credentials", return_value=None),
                        patch.object(identity_storage, "_load_credentials", return_value=None),
                        patch.object(tool_auth, "authenticated_identity", new=self.identity),
                        patch.object(generated_storage, "oss_put_editor_revision", new=self.writer),
                        patch.object(generated_storage, "read_generated_bytes", new=AsyncMock(side_effect=self.read)),
                        patch.object(personal_service, "_user_permission_payload", new=AsyncMock(return_value={
                            "menu_keys": ["task"], "action_keys": ["task.read"], "data_scope": "本人数据",
                        })),
                        patch.object(personal_service.settings, "personal_agent_root", str(root / "runtime")),
                        patch.object(personal_service.settings, "ai_oss_bucket", "codex-ai-test"),
                        patch.object(personal_service.settings, "ai_oss_prefix", "CODEX/AI空间"),
                        patch.object(personal_service.settings, "ai_oss_endpoint", "https://oss.invalid"),
                        patch.object(personal_service.settings, "ai_oss_credentials_file", "isolated-credentials.csv")]
        for item in self.patches:
            item.start()
        self.context = tool_auth.AuthContext(None, "isolated-request-token", {
            "method": "POST", "path": personal_service.settings.api_prefix + "/personal-agent/messages",
        })
        self.context_token = tool_auth._auth_context.set(self.context)

    async def asyncTearDown(self):
        tool_auth._auth_context.reset(self.context_token)
        for item in reversed(self.patches):
            item.stop()
        await self.db.close()
        await self.engine.dispose()
        self.temporary.cleanup()

    async def identity(self, db=None):
        return {"username": self.username}

    async def put(self, bucket, key, content, content_type, *, endpoint):
        self.objects[f"oss://{bucket}/{key}"] = content
        return True

    async def read(self, item, *, max_bytes):
        return self.objects[item.path]

    async def safe_identity(self):
        return await personal_service.personal_identity(await self.identity(), self.db)

    async def persist(self, *, db=None, conversation_id="default"):
        identity = await self.safe_identity()
        return await identity_storage.persist_personal_identity(
            db or self.db, identity, content=personal_service._identity_markdown(identity), conversation_id=conversation_id,
        )

    async def claims(self):
        return (await self.db.scalars(select(BusinessRecord).where(
            BusinessRecord.module == SPACE_MODULE, BusinessRecord.serial_no.like("PAID-%"),
        ).execution_options(populate_existing=True))).all()

    async def test_post_archives_only_actual_identity_before_model_and_preserves_local_file(self):
        forged = {"username": self.actor.username, "display_name": "伪造姓名", "role": "admin", "action_keys": ["all"]}

        async def generate(identity, db, text, on_delta, **kwargs):
            self.assertEqual(identity["role"], "user")
            self.assertEqual(identity["action_keys"], ["task.read"])
            self.assertEqual(len(self.objects), 1)
            expected = personal_service._identity_markdown(identity)
            self.assertEqual(next(iter(self.objects.values())), expected.encode("utf-8"))
            self.assertEqual(personal_service._identity_path(self.actor.username).read_text(encoding="utf-8"), expected)
            return {"state": kwargs["state"], "response": "模型回答不是身份文件来源"}

        with patch.object(personal_service, "_generate_response", new=AsyncMock(side_effect=generate)):
            await personal_service.generate_response(forged, self.db, "第一次真实提问", AsyncMock())
            await personal_service.generate_response(forged, self.db, "再次提问", AsyncMock())
        self.assertEqual(self.writer.await_count, 1)
        files = await list_space_files(await self.identity(), self.db, page=1, page_size=20)
        self.assertEqual(files["total"], 1)
        self.assertEqual(files["items"][0]["source_kind"], "personal_identity")
        self.assertEqual(files["items"][0]["name"], "IDENTITY.md")
        self.assertEqual(files["items"][0]["conversation_id"], "default")

    async def test_reads_never_claim_or_put_identity(self):
        self.context.scope["method"] = "GET"
        identity = await self.safe_identity()
        await personal_service.personal_state(identity, self.db)
        await personal_service.manage_conversations(identity, db=self.db)
        self.assertEqual(await self.claims(), [])
        self.writer.assert_not_awaited()

    async def test_versions_reuse_history_hash_and_keep_original_bytes(self):
        first = await self.persist()
        original = dict(self.objects)
        self.actor.display_name = "真实改名员工"
        await self.db.commit()
        second = await self.persist(conversation_id="other-conversation")
        self.assertNotEqual(first["id"], second["id"])
        self.assertNotEqual(first["sha256"], second["sha256"])
        self.actor.display_name = "真实员工"
        await self.db.commit()
        again = await self.persist(conversation_id="third-conversation")
        self.assertEqual(first["id"], again["id"])
        self.assertEqual(self.writer.await_count, 2)
        self.assertTrue(all(self.objects[path] == raw for path, raw in original.items()))
        self.assertEqual(len(await self.claims()), 2)
        self.assertTrue(all(len(row.serial_no) <= 64 and row.status == "available" for row in await self.claims()))

    async def test_real_auth_and_post_origin_are_required(self):
        identity = await self.safe_identity()
        for method, path in (("GET", "/personal-agent/state"), ("POST", "/personal-agent/conversations")):
            self.context.scope.update(method=method, path=personal_service.settings.api_prefix + path)
            with self.assertRaises(HTTPException) as raised:
                await self.persist()
            self.assertEqual(raised.exception.status_code, 403)
        self.context.scope.update(method="POST", path=personal_service.settings.api_prefix + "/personal-agent/messages")
        with self.assertRaises(HTTPException) as raised:
            await identity_storage.persist_personal_identity(
                self.db, {**identity, "username": "another-person"}, content="不可信内容", conversation_id="default")
        self.assertEqual(raised.exception.status_code, 403)
        self.context.active = False
        with self.assertRaises(HTTPException) as raised:
            await self.persist()
        self.assertEqual(raised.exception.status_code, 401)
        self.writer.assert_not_awaited()
        self.assertEqual(await self.claims(), [])

    async def test_failed_and_uncertain_claims_never_automatically_retry(self):
        with patch.object(identity_storage, "save_generated_file", new=AsyncMock(side_effect=HTTPException(503, "配置失败"))):
            with self.assertRaises(HTTPException):
                await self.persist()
        self.assertEqual((await self.claims())[0].status, "failed")
        with self.assertRaises(HTTPException) as raised:
            await self.persist(conversation_id="different-conversation")
        self.assertEqual(raised.exception.detail["status"], "failed")
        self.writer.assert_not_awaited()
        self.actor.display_name = "新内容员工"
        await self.db.commit()
        with patch.object(generated_storage, "read_generated_bytes", new=AsyncMock(side_effect=HTTPException(502, "读取失败"))):
            with self.assertRaises(HTTPException):
                await self.persist()
        with self.assertRaises(HTTPException) as raised:
            await self.persist()
        self.assertEqual(raised.exception.detail["status"], "uncertain")
        self.assertEqual(self.writer.await_count, 1)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)

    async def test_cancelled_upload_retains_uncertain_claim(self):
        with patch.object(generated_storage, "oss_put_editor_revision", new=AsyncMock(side_effect=asyncio.CancelledError)):
            with self.assertRaises(asyncio.CancelledError):
                await self.persist()
        self.assertEqual((await self.claims())[0].status, "uncertain")
        with self.assertRaises(HTTPException) as raised:
            await self.persist()
        self.assertEqual(raised.exception.detail["status"], "uncertain")
        self.writer.assert_not_awaited()

    async def test_claim_available_is_atomic_with_file_commit_and_lost_ack_does_not_reupload(self):
        commit = self.db.commit
        count = 0

        async def lost_ack():
            nonlocal count
            count += 1
            await commit()
            if count == 3:
                raise RuntimeError("isolated committed acknowledgment lost")

        with patch.object(self.db, "commit", new=lost_ack):
            with self.assertRaises(HTTPException):
                await self.persist()
        claim = (await self.claims())[0]
        self.assertEqual(claim.status, "available")
        file = await self.db.get(BusinessRecord, claim.data["file_record_id"])
        self.assertEqual(file.status, "available")
        self.assertEqual(file.data["file_id"], claim.data["file_id"])
        again = await self.persist()
        self.assertEqual(again["id"], claim.data["file_id"])
        self.assertEqual(self.writer.await_count, 1)

    async def test_final_commit_failure_does_not_publish_claim_or_replay(self):
        commit = self.db.commit
        count = 0

        async def fail_final():
            nonlocal count
            count += 1
            if count == 3:
                raise RuntimeError("isolated final transaction failed")
            await commit()

        with patch.object(self.db, "commit", new=fail_final):
            with self.assertRaises(HTTPException):
                await self.persist()
        claim = (await self.claims())[0]
        self.assertEqual(claim.status, "uncertain")
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 0)
        with self.assertRaises(HTTPException):
            await self.persist()
        self.assertEqual(self.writer.await_count, 1)

    async def test_different_sessions_compete_on_persistent_unique_claim(self):
        identity = await self.safe_identity()
        await generated_storage.resolve_generated_scope(self.db, identity)
        await self.db.commit()
        original = identity_storage._claim
        arrivals = 0
        ready = asyncio.Event()

        async def simultaneous(db, serial):
            nonlocal arrivals
            claim = await original(db, serial)
            if claim is None:
                arrivals += 1
                if arrivals == 2:
                    ready.set()
                await asyncio.wait_for(ready.wait(), 5)
            return claim

        async def contender(conversation_id):
            async with self.sessions() as db:
                return await identity_storage.persist_personal_identity(
                    db, identity, content=personal_service._identity_markdown(identity), conversation_id=conversation_id)

        with patch.object(identity_storage, "_claim", new=simultaneous):
            results = await asyncio.gather(contender("first-conversation"), contender("second-conversation"), return_exceptions=True)
        for result in results:
            if isinstance(result, HTTPException):
                self.assertEqual(result.status_code, 409)
            else:
                self.assertIsInstance(result, dict)
                self.assertEqual(result["status"], "available")
        self.assertTrue(any(isinstance(result, dict) for result in results))
        self.assertEqual(self.writer.await_count, 1)
        self.assertEqual(len(await self.claims()), 1)
        self.assertEqual(await self.db.scalar(select(func.count()).select_from(FileAttachment)), 1)

    async def test_uploading_claim_cannot_be_taken_over_by_another_session(self):
        identity = await self.safe_identity()
        started = asyncio.Event()
        finish = asyncio.Event()

        async def blocked_put(*args, **kwargs):
            started.set()
            await asyncio.wait_for(finish.wait(), 5)
            return await self.put(*args, **kwargs)

        with patch.object(generated_storage, "oss_put_editor_revision", new=AsyncMock(side_effect=blocked_put)) as writer:
            first = asyncio.create_task(self.persist())
            try:
                await asyncio.wait_for(started.wait(), 5)
                async with self.sessions() as db:
                    with self.assertRaises(HTTPException) as raised:
                        await identity_storage.persist_personal_identity(
                            db, identity, content=personal_service._identity_markdown(identity), conversation_id="other-session")
                    self.assertEqual(raised.exception.detail["status"], "uploading")
                    self.assertEqual(raised.exception.status_code, 409)
            finally:
                finish.set()
                await first
            self.assertEqual(writer.await_count, 1)
        self.assertEqual((await self.claims())[0].status, "available")


class PersonalConfirmationBoundaryTest(IsolatedAsyncioTestCase):
    async def test_personal_office_routes_preserve_frontend_confirmation_schema_and_result_contract(self):
        from fastapi import FastAPI
        from httpx import ASGITransport, AsyncClient
        from app.areas.personal_agent import ai_space, dingtalk_office_service

        application = FastAPI()
        application.include_router(ai_space.router)
        identity = {"username": "CODEX-office-route-owner"}
        application.dependency_overrides[ai_space.current_identity] = lambda: identity
        application.dependency_overrides[ai_space.get_db] = lambda: None
        file_id, sha256, session_id = "a" * 32, "b" * 64, "c" * 32
        valid = {"Origin": "http://localhost", "X-OA-Agent-Confirmation": "frontend"}
        opened = {"session_id": session_id, "name": "原生.docx", "url": "https://alidocs.dingtalk.com/CODEX"}
        saved = {"file": {"id": "d" * 32, "supersedes": file_id, "can_delete": True}, "supersedes": file_id, "unchanged": False}
        with patch.object(dingtalk_office_service, "start_office_edit", new=AsyncMock(return_value=opened)) as open_file, \
                patch.object(dingtalk_office_service, "save_office_edit", new=AsyncMock(return_value=saved)) as save_file:
            async with AsyncClient(transport=ASGITransport(app=application), base_url="http://isolated.invalid") as client:
                for route, body, expected, status, operation, argument in (
                    ("open", {"expected_sha256": sha256}, opened, 200, open_file, sha256),
                    ("save", {"session_id": session_id}, saved, 201, save_file, session_id),
                ):
                    path = personal_service.settings.api_prefix + f"/personal-agent/ai-space/files/{file_id}/office/{route}"
                    for headers in ({}, {"Origin": "http://localhost"}, {**valid, "Origin": "https://untrusted.invalid"}):
                        response = await client.post(path, json=body, headers=headers)
                        self.assertEqual(response.status_code, 403)
                    for invalid in ({}, {key: "invalid" for key in body}, {**body, "owner_user_id": 1}):
                        response = await client.post(path, json=invalid, headers=valid)
                        self.assertEqual(response.status_code, 422)
                    operation.assert_not_awaited()
                    response = await client.post(path, json=body, headers=valid)
                    self.assertEqual(response.status_code, status)
                    self.assertEqual(response.json(), expected)
            operation.assert_awaited_once_with(file_id, argument, identity, None, create_copy=False)

    def test_office_provider_resources_keep_https_domain_identifier_and_header_boundaries(self):
        from app.areas.personal_agent import dingtalk_office_client as office

        self.assertEqual(office.provider_url("https://alidocs.dingtalk.com/CODEX", editor=True), "https://alidocs.dingtalk.com/CODEX")
        for url in ("http://dingtalk.com/CODEX", "https://dingtalk.com.evil.invalid/CODEX",
                    "https://account@dingtalk.com/CODEX", "https://dingtalk.com:444/CODEX", "https://aliyuncs.com/CODEX"):
            with self.subTest(url=url), self.assertRaises(HTTPException) as raised:
                office.provider_url(url, editor=True)
            self.assertEqual(raised.exception.status_code, 502)
        for identifier in ("", "a/b", "a\\b", "a?b", "a\nb", "a" * 257):
            with self.subTest(identifier=identifier), self.assertRaises(HTTPException):
                office.identifier(identifier)
        body = {"protocol": "HEADER_SIGNATURE", "headerSignatureInfo": {
            "resourceUrls": ["https://storage.aliyuncs.com/CODEX"], "headers": {"x-codex": "isolated"},
        }}
        self.assertEqual(office._signed_resource(body), ("https://storage.aliyuncs.com/CODEX", {"x-codex": "isolated"}))
        body["headerSignatureInfo"]["headers"]["x-codex"] = "value\r\ninjected"
        with self.assertRaises(HTTPException) as raised:
            office._signed_resource(body)
        self.assertEqual(raised.exception.status_code, 502)

    async def test_personal_file_delete_route_requires_exact_hash_trusted_origin_and_frontend_confirmation(self):
        from fastapi import FastAPI
        from httpx import ASGITransport, AsyncClient
        from app.areas.personal_agent import ai_space, ai_space_deletion

        application = FastAPI()
        application.include_router(ai_space.router)
        application.dependency_overrides[ai_space.current_identity] = lambda: {"username": "CODEX-route-owner"}
        application.dependency_overrides[ai_space.get_db] = lambda: None
        file_id, sha256 = "a" * 32, "b" * 64
        path = personal_service.settings.api_prefix + f"/personal-agent/ai-space/files/{file_id}"
        valid = {"Origin": "http://localhost", "X-OA-Agent-Confirmation": "frontend"}
        result = {"id": file_id, "status": "deleted", "blob_status": "absent"}
        with patch.object(ai_space_deletion, "delete_space_file", new=AsyncMock(return_value=result)) as delete:
            async with AsyncClient(transport=ASGITransport(app=application), base_url="http://isolated.invalid") as client:
                for headers in ({}, {"Origin": "http://localhost"}, {**valid, "Origin": "https://untrusted.invalid"}):
                    response = await client.request("DELETE", path, json={"expected_sha256": sha256}, headers=headers)
                    self.assertEqual(response.status_code, 403)
                delete.assert_not_awaited()
                for body in ({}, {"expected_sha256": "wrong"}, {"expected_sha256": sha256, "owner_user_id": 1}):
                    response = await client.request("DELETE", path, json=body, headers=valid)
                    self.assertEqual(response.status_code, 422)
                delete.assert_not_awaited()
                response = await client.request("DELETE", path, json={"expected_sha256": sha256}, headers=valid)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json(), result)
                delete.assert_awaited_once_with(file_id, sha256, {"username": "CODEX-route-owner"}, None)

    def request(self, status, owner="本人"):
        return BusinessRecord(
            id=1, module="agent_mcp_action", serial_no="MCP-request-1", owner=owner,
            status=status, data={"request_id": "request-1", "error": {"status_code": 409, "detail": {"detail": "原业务拒绝"}} if status == "执行失败" else None},
        )

    def state(self):
        return {"conversation_id": "default", "messages": [], "pending_actions": [{
            "id": "action-1", "type": "mcp.call", "status": "pending", "requested_by": "本人",
            "payload": {"request_id": "request-1"}, "summary": "旧待确认卡",
        }]}

    async def test_persistent_status_projection_is_read_only_and_identity_scoped(self):
        for stored, visible in (("执行失败", "failed"), ("已执行", "succeeded"), ("已拒绝", "rejected"), ("执行中", "executing"), ("待确认", "pending")):
            with self.subTest(status=stored):
                state = self.state()
                original = deepcopy(state)
                record = self.request(stored)
                db = SimpleNamespace(scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [record])))
                result = await project_personal_actions(state, {"username": "本人"}, db)
                self.assertEqual(result["pending_actions"][0]["status"], visible)
                self.assertEqual(state, original)
                self.assertEqual(record.status, stored)
                self.assertIsNot(result["pending_actions"][0], state["pending_actions"][0])
                query = db.scalars.await_args.args[0]
                parameters = query.compile().params
                self.assertIn("本人", parameters.values())
                self.assertIn("agent_mcp_action", parameters.values())
                self.assertIn(["MCP-request-1"], parameters.values())
                if stored == "执行失败":
                    self.assertEqual(result["pending_actions"][0]["error"]["detail"]["detail"], "原业务拒绝")
        db = SimpleNamespace(scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [])))
        with self.assertRaises(HTTPException) as raised:
            await project_personal_actions(self.state(), {"username": "本人"}, db)
        self.assertEqual(raised.exception.status_code, 404, "其他所有者的请求不能被投影")
        foreign = self.state()
        foreign["pending_actions"][0]["requested_by"] = "别人"
        db.scalars.reset_mock()
        with self.assertRaises(HTTPException) as raised:
            await project_personal_actions(foreign, {"username": "本人"}, db)
        self.assertEqual(raised.exception.status_code, 403)
        db.scalars.assert_not_awaited()
        for action_type in ("create_task", "approve_contract"):
            for stored, visible in (("执行失败", "failed"), ("已执行", "succeeded"), ("已拒绝", "rejected"), ("执行中", "executing"), ("待确认", "pending")):
                with self.subTest(action_type=action_type, status=stored):
                    state = self.state()
                    state["pending_actions"][0]["type"] = action_type
                    original = deepcopy(state)
                    record = self.request(stored)
                    record.module, record.serial_no = "personal_agent_action", "PA-action-1"
                    db = SimpleNamespace(scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [record])))
                    result = await project_personal_actions(state, {"username": "本人"}, db)
                    self.assertEqual(result["pending_actions"][0]["status"], visible)
                    self.assertEqual(result["pending_actions"][0]["error"], record.data["error"])
                    self.assertEqual(state, original)
                    self.assertEqual(record.status, stored)
            db = SimpleNamespace(scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [])))
            result = await project_personal_actions(state, {"username": "本人"}, db)
            self.assertEqual(result["pending_actions"][0], state["pending_actions"][0], "升级前无 PA 存档时保留原 pending")
            db = SimpleNamespace(scalars=AsyncMock(side_effect=[SimpleNamespace(all=lambda: []), SimpleNamespace(all=lambda: ["PA-action-1"])]))
            with self.assertRaises(HTTPException) as raised:
                await project_personal_actions(state, {"username": "本人"}, db)
            self.assertEqual(raised.exception.status_code, 403, "他人的 PA 存档不能冒充无存档旧请求")

    async def test_personal_state_and_conversation_reads_project_without_saving(self):
        workspace = {"kind": "personal"}
        state = {**self.state(), "workspace": workspace}
        original = deepcopy(state)
        record = self.request("执行失败")
        db = SimpleNamespace(scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [record])))
        with patch.object(personal_service, "ensure_identity_file", new=AsyncMock(return_value="身份")), \
             patch.object(personal_service, "_read_state", return_value=state), \
             patch.object(personal_service, "_user_dir", return_value=Path("isolated-unit")), \
             patch.object(personal_service.conversations, "list_conversations", new=AsyncMock(return_value={
                 "state": state, "items": [], "active_id": "default", "workspace": workspace, "workspaces": [workspace],
             })), \
             patch.object(personal_service.conversations, "save_state", new=AsyncMock()) as save:
            identity = {"username": "本人"}
            direct = await personal_service.personal_state(identity, db)
            snapshot = await personal_service.manage_conversations(identity, db=db)
            self.assertEqual(direct["pending_actions"][0]["status"], "failed")
            self.assertEqual(snapshot["state"]["pending_actions"][0]["status"], "failed")
            self.assertEqual(snapshot["state"]["workspace"], workspace)
            self.assertEqual(snapshot["workspaces"], [{**workspace, "title": "个人空间", "available": True}])
            self.assertEqual(state, original)
            save.assert_not_awaited()

    async def test_decided_mcp_requests_remain_non_replayable(self):
        for status in ("执行失败", "已执行", "已拒绝", "执行中"):
            with self.subTest(status=status):
                record = self.request(status)
                db = SimpleNamespace(scalar=AsyncMock(return_value=record))
                with patch.object(tool_service, "require_human_decision"), \
                     patch.object(tool_service, "authenticated_identity", new=AsyncMock(return_value={"username": "本人"})), \
                     patch.object(tool_service, "execute_original", new=AsyncMock()) as execute:
                    with self.assertRaises(HTTPException) as raised:
                        await tool_service.decide_tool_request("request-1", "approved", {"username": "本人"}, db)
                    self.assertEqual(raised.exception.status_code, 409)
                    self.assertEqual(record.status, status)
                    execute.assert_not_awaited()
                for action_type in ("create_task", "approve_contract"):
                    action = {**self.state()["pending_actions"][0], "type": action_type}
                    record.module, record.serial_no = "personal_agent_action", "PA-action-1"
                    with self.assertRaises(HTTPException) as raised:
                        await begin_personal_decision(action, "approved", "", {"username": "本人"}, self.state(), db)
                    self.assertEqual(raised.exception.status_code, 409)

    async def test_action_dismissal_persists_only_presentation_and_keeps_audit(self):
        with TemporaryDirectory(prefix="oa-test-action-dismiss-") as temporary:
            directory = Path(temporary)
            snapshot = await conversations.create_conversation(directory)
            state = {**snapshot["state"], "pending_actions": self.state()["pending_actions"]}
            await conversations.save_state(directory, state)
            original = conversations.read_state(directory, state["conversation_id"])
            record = self.request("执行失败")
            db = SimpleNamespace(scalar=AsyncMock(return_value="本人"),
                                 scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [record])), commit=AsyncMock())
            result = await dismiss_personal_actions(directory, state["conversation_id"], ["action-1"], {"username": "本人"}, db)
            stored = conversations.read_state(directory, state["conversation_id"])
            self.assertEqual(stored["dismissed_action_ids"], ["action-1"])
            self.assertEqual({**stored, "dismissed_action_ids": []}, original)
            self.assertEqual(result["state"]["pending_actions"][0]["status"], "failed")
            self.assertEqual(stored["pending_actions"][0]["status"], "pending", "原结构化请求不被展示投影重写")
            self.assertEqual(record.status, "执行失败")
            with patch.object(conversations, "_write_json", wraps=conversations._write_json) as write:
                again = await dismiss_personal_actions(directory, state["conversation_id"], ["action-1"], {"username": "本人"}, db)
                self.assertEqual(again["state"]["dismissed_action_ids"], ["action-1"])
                write.assert_not_called()
            await conversations.save_state(directory, {**stored, "dismissed_action_ids": ["departed-action", "action-1"]})
            with patch.object(conversations, "_write_json", wraps=conversations._write_json) as write:
                self.assertEqual(conversations.read_state(directory, state["conversation_id"])["dismissed_action_ids"],
                                 ["departed-action", "action-1"])
                write.assert_not_called()
                cleaned = await dismiss_personal_actions(directory, state["conversation_id"], ["action-1"], {"username": "本人"}, db)
                self.assertEqual(cleaned["state"]["dismissed_action_ids"], ["action-1"])
                self.assertEqual(conversations.read_state(directory, state["conversation_id"]), stored)
                self.assertEqual(write.call_count, 1)
            await conversations.save_state(directory, stored)
            self.assertEqual(conversations.read_state(directory, state["conversation_id"])["dismissed_action_ids"], ["action-1"])
            db.commit.assert_not_awaited()
            async with conversations.conversation_operation(directory, state["conversation_id"]):
                with self.assertRaises(HTTPException) as raised:
                    await dismiss_personal_actions(directory, state["conversation_id"], ["action-1"], {"username": "本人"}, db)
                self.assertEqual(raised.exception.status_code, 409)
            record.owner = "别人"
            with self.assertRaises(HTTPException) as raised:
                await dismiss_personal_actions(directory, state["conversation_id"], ["action-1"], {"username": "本人"}, db)
            self.assertEqual(raised.exception.status_code, 403)
            self.assertEqual(conversations.read_state(directory, state["conversation_id"]), stored)

    async def test_batch_dismissal_validates_all_ids_before_one_presentation_write(self):
        with TemporaryDirectory(prefix="oa-test-action-batch-dismiss-") as temporary:
            directory = Path(temporary)
            snapshot = await conversations.create_conversation(directory)
            first = self.state()["pending_actions"][0]
            second = {**first, "id": "action-2", "payload": {"request_id": "request-2"}}
            state = {**snapshot["state"], "pending_actions": [first, second],
                     "dismissed_action_ids": ["departed-action", "action-1"]}
            await conversations.save_state(directory, state)
            record = self.request("待确认")
            other = self.request("待确认")
            other.id, other.serial_no = 2, "MCP-request-2"
            db = SimpleNamespace(scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: [record, other])), commit=AsyncMock())
            before = conversations.read_state(directory, state["conversation_id"])
            with patch.object(conversations, "_write_json", wraps=conversations._write_json) as write:
                self.assertEqual(conversations.read_state(directory, state["conversation_id"])["dismissed_action_ids"],
                                 ["departed-action", "action-1"])
                write.assert_not_called()
                for ids, expected in ((["action-1", "missing"], 404), (["action-1", "action-1"], 422), ([], 422)):
                    with self.assertRaises(HTTPException) as raised:
                        await dismiss_personal_actions(directory, state["conversation_id"], ids, {"username": "本人"}, db)
                    self.assertEqual(raised.exception.status_code, expected)
                    self.assertEqual(conversations.read_state(directory, state["conversation_id"]), before)
                other.owner = "别人"
                with self.assertRaises(HTTPException) as raised:
                    await dismiss_personal_actions(directory, state["conversation_id"], ["action-1", "action-2"], {"username": "本人"}, db)
                self.assertEqual(raised.exception.status_code, 403)
                write.assert_not_called()
                other.owner = "本人"
                result = await dismiss_personal_actions(directory, state["conversation_id"], ["action-1", "action-2"], {"username": "本人"}, db)
                self.assertEqual(write.call_count, 1)
                self.assertEqual(result["state"]["dismissed_action_ids"], ["action-1", "action-2"])
                self.assertEqual({**conversations.read_state(directory, state["conversation_id"]),
                                  "dismissed_action_ids": before["dismissed_action_ids"]}, before)
                self.assertEqual(len(result["state"]["pending_actions"]), 2)
                self.assertEqual(record.status, "待确认")
                self.assertEqual(other.status, "待确认")
                db.commit.assert_not_awaited()

    async def test_pending_history_preserves_legacy_and_real_conversation_source(self):
        with TemporaryDirectory(prefix="oa-test-action-history-") as temporary:
            directory = Path(temporary)
            snapshot = await conversations.create_conversation(directory)
            original = self.state()["pending_actions"][0]
            original["preview"] = {"operation_name": "真实操作", "method": "POST", "path_template": "/api/v1/isolated-operation",
                                   "params": {"path": {}, "query": {}, "body": {}}}
            legacy = {"id": "legacy-action", "type": "create_task", "requested_by": "本人", "status": "pending",
                      "payload": {"title": "旧任务"}, "summary": "旧任务", "created_at": "2026-10-13T00:00:00+00:00"}
            state = {**snapshot["state"], "pending_actions": [original, legacy]}
            await conversations.save_state(directory, state)
            before = conversations.read_state(directory, state["conversation_id"])
            record = self.request("待确认")
            record.data = {**record.data, "tool_name": "真实操作", "summary": "真实操作", "preview": original["preview"]}
            db = SimpleNamespace(scalar=AsyncMock(return_value=1),
                                 scalars=AsyncMock(side_effect=[SimpleNamespace(all=lambda: []),
                                                              SimpleNamespace(all=lambda: [record]), SimpleNamespace(all=lambda: [])]),
                                 commit=AsyncMock())
            history = await read_confirmation_history({"username": "本人"}, db, status="pending", directory=directory)
            self.assertEqual(history["total"], 2)
            self.assertEqual([item["status"] for item in history["items"]], ["pending", "pending"])
            self.assertTrue(all(item["can_decide"] for item in history["items"]))
            self.assertTrue(all(item["conversation_id"] == state["conversation_id"] for item in history["items"]))
            self.assertEqual(history["items"][0]["id"], "request-1")
            self.assertEqual(history["items"][0]["action_id"], "action-1")
            db.scalars = AsyncMock(side_effect=[SimpleNamespace(all=lambda: []), SimpleNamespace(all=lambda: [])])
            second = await read_confirmation_history({"username": "本人"}, db, status="pending", page=2, page_size=1, directory=directory)
            self.assertEqual(second["total"], 2)
            self.assertEqual([item["id"] for item in second["items"]], ["legacy-action"])
            record.data["context"] = {"conversation_id": state["conversation_id"], "conversation_title": state["conversation_title"]}
            for stored, visible in (("待确认", "pending"), ("执行失败", "failed"), ("已执行", "succeeded"), ("已拒绝", "rejected")):
                record.status = stored
                db.scalar = AsyncMock(side_effect=[record, None])
                detail = await read_confirmation_detail("mcp", "request-1", {"username": "本人"}, db, directory=directory)
                self.assertEqual(detail["action"]["id"], "action-1")
                self.assertEqual(detail["action"]["status"], visible)
                self.assertEqual(detail["can_decide"], visible == "pending")
            db.scalar = AsyncMock(return_value=None)
            detail = await read_confirmation_detail("personal_action", "legacy-action", {"username": "本人"}, db, directory=directory)
            self.assertTrue(detail["can_decide"])
            self.assertEqual(detail["action"]["payload"], legacy["payload"])
            record.owner = "别人"
            db.scalar = AsyncMock(return_value=record)
            with self.assertRaises(HTTPException) as raised:
                await read_confirmation_detail("mcp", "request-1", {"username": "本人"}, db, directory=directory)
            self.assertEqual(raised.exception.status_code, 403)
            self.assertEqual(conversations.read_state(directory, state["conversation_id"]), before)
            db.commit.assert_not_awaited()

    async def test_truncated_mcp_card_keeps_true_preview_without_fabricating_action(self):
        with TemporaryDirectory(prefix="oa-test-action-truncated-") as temporary:
            directory = Path(temporary)
            snapshot = await conversations.create_conversation(directory)
            record = self.request("待确认")
            record.data = {**record.data, "tool_name": "原工具", "summary": "原请求", "preview": {"params": {"body": {"value": "原参数"}}},
                           "context": {"conversation_id": snapshot["state"]["conversation_id"], "conversation_title": "原对话"}}
            db = SimpleNamespace(scalar=AsyncMock(side_effect=[record, None]), commit=AsyncMock())
            detail = await read_confirmation_detail("mcp", "request-1", {"username": "本人"}, db, directory=directory)
            self.assertEqual(detail["conversation_id"], snapshot["state"]["conversation_id"])
            self.assertEqual(detail["request_preview"], record.data["preview"])
            self.assertIsNone(detail["action"])
            self.assertIsNone(detail["item"]["action_id"])
            self.assertFalse(detail["can_decide"], "30条截断后不能捏造原操作ID或自动恢复可执行卡")
            self.assertEqual(conversations.read_state(directory)["pending_actions"], [])
            db.commit.assert_not_awaited()

    async def test_removed_personal_actions_use_archive_without_becoming_replayable(self):
        with TemporaryDirectory(prefix="oa-test-action-archive-") as temporary:
            directory = Path(temporary)
            record = BusinessRecord(module="personal_agent_action", owner="本人", status="执行失败", title="原操作",
                                    data={"action_id": "archived-action", "action_type": "create_task", "summary": "原操作",
                                          "payload": {"title": "原任务"}, "conversation_id": "default", "conversation_title": "原对话",
                                          "error": {"status_code": 409, "detail": "原业务拒绝"}})
            db = SimpleNamespace(scalar=AsyncMock(side_effect=[record, None]), commit=AsyncMock())
            detail = await read_confirmation_detail("personal_action", "archived-action", {"username": "本人"}, db, directory=directory)
            self.assertEqual(detail["action"]["payload"], record.data["payload"])
            self.assertEqual(detail["action"]["status"], "failed")
            self.assertEqual(detail["item"]["error"], record.data["error"])
            self.assertEqual(detail["conversation_id"], "default")
            self.assertFalse(detail["can_decide"])
            self.assertEqual(conversations.read_state(directory)["pending_actions"], [])
            self.assertEqual(list(directory.iterdir()), [], "读取历史不能生成会话或补建业务存档")
            db.commit.assert_not_awaited()


class DomainBoundaryTest(IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.tempdir = TemporaryDirectory(prefix="oa-test-enterprise-boundaries-")
        self.upload_root = Path(self.tempdir.name) / "uploads"
        self.upload_root.mkdir()
        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False, class_=AsyncSession)
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)

    async def asyncTearDown(self) -> None:
        await self.engine.dispose()
        self.tempdir.cleanup()

    async def test_settlement_guard_keeps_legacy_inactive_statuses(self) -> None:
        async with self.sessions() as db:
            rows = [BusinessRecord(
                module="finance_settlement", serial_no=f"ENTERPRISE-BOUNDARIES-{index}",
                title="隔离结算", customer="隔离客户", status=record_status,
                owner="enterprise-admin", data={"receipt_id": "42"},
            ) for index, record_status in enumerate(
                ["已拒绝", "已驳回", "已退回", "已撤回", "已作废", "待审批"], start=1,
            )]
            db.add_all(rows)
            await db.commit()
            active = await _active_settlements_by_receipt(db, {42})
            self.assertEqual(active[42].id, rows[-1].id)
            self.assertEqual(
                await _active_settlements_by_receipt(db, {42}, exclude_application_ids={rows[-1].id}),
                {},
            )

    async def test_generic_case_import_remains_blocked(self) -> None:
        file = UploadFile(file=BytesIO(b"serial_no,title\nCASE-1,Case 1\n"), filename="cases.csv")
        with patch("app.core.permissions._require_record_module_menu", new=AsyncMock()):
            with self.assertRaises(HTTPException) as raised:
                await import_business_records(
                    module="case", file=file,
                    identity={"username": "enterprise-admin", "role": "admin"}, db=object(),
                )
        self.assertEqual(raised.exception.status_code, 409)

    async def test_allocation_reversal_remains_in_callers_transaction(self) -> None:
        async with self.sessions() as db:
            contract = BusinessRecord(
                module="contract", serial_no="ENTERPRISE-BOUNDARIES-CONTRACT",
                title="隔离合同", customer="隔离客户", status="已审批", owner="enterprise-admin", data={},
            )
            fee = BusinessRecord(
                module="finance", serial_no="ENTERPRISE-BOUNDARIES-FEE",
                title="隔离费用", customer="隔离客户", status="已收款", owner="enterprise-admin",
                data={"received_amount": 100, "cashed_amount": 100, "incoming_payment_id": 42},
            )
            db.add_all([contract, fee])
            await db.flush()
            plan = ReceivablePlan(
                contract_record_id=contract.id, phase="代理费", due_date=date.today(),
                amount=100, received_amount=100, status="已收款",
            )
            transaction = FinanceTransaction(
                finance_record_id=fee.id, transaction_type="收款", amount=100,
                transaction_date=date.today(), operator="enterprise-admin",
            )
            db.add_all([plan, transaction])
            await db.commit()
            fee_id, plan_id, transaction_id = fee.id, plan.id, transaction.id

        async with self.sessions() as db:
            await _revert_incoming_allocation(
                {"amount": 100, "receivable_plan_id": plan_id, "fee_record_id": fee_id,
                 "transaction_id": transaction_id},
                db, payment_id=42,
            )
            await db.flush()
            self.assertEqual((await db.get(ReceivablePlan, plan_id)).received_amount, 0)
            self.assertIsNone(await db.get(FinanceTransaction, transaction_id))
            await db.rollback()

        async with self.sessions() as db:
            self.assertEqual((await db.get(ReceivablePlan, plan_id)).received_amount, 100)
            self.assertEqual((await db.get(BusinessRecord, fee_id)).data["received_amount"], 100)
            self.assertIsNotNone(await db.get(FinanceTransaction, transaction_id))

    async def test_seal_file_and_row_restore_when_commit_fails(self) -> None:
        identity = {"username": "enterprise-admin", "role": "admin", "department": "测试"}
        file_path = self.upload_root / "seal-proof.pdf"
        file_path.write_bytes(b"seal-proof")
        async with self.sessions() as db:
            db.add(User(
                username=identity["username"], display_name="测试管理员", role="admin",
                department="测试", password_hash="unused", is_active=True,
            ))
            record = BusinessRecord(
                module="seal", serial_no="ENTERPRISE-BOUNDARIES-SEAL",
                title="隔离用印", customer="隔离客户", status="草稿",
                owner=identity["username"], department="测试", data={"document_names": file_path.name},
            )
            db.add(record)
            await db.flush()
            attachment = FileAttachment(
                record_id=record.id, category="用印文件", original_name=file_path.name,
                stored_name=file_path.name, path=str(file_path), content_type="application/pdf",
                size=file_path.stat().st_size, uploader=identity["username"],
            )
            db.add(attachment)
            await db.commit()
            attachment_id = attachment.id

            with patch.object(seals, "UPLOAD_ROOT", self.upload_root):
                with patch.object(db, "commit", new=AsyncMock(side_effect=RuntimeError("commit failed"))):
                    with self.assertRaisesRegex(RuntimeError, "commit failed"):
                        await seals.batch_delete_seal_attachments(
                            AttachmentBatchInput(attachment_ids=[attachment_id]), identity, db,
                        )
            self.assertIsNotNone(await db.get(FileAttachment, attachment_id))
            self.assertEqual(file_path.read_bytes(), b"seal-proof")
            self.assertFalse(list(self.upload_root.glob(".pending-delete-*")))
