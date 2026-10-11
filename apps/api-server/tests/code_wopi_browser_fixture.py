"""真实 CODE 浏览器验收的隔离初始化、只读回执与精确清理入口。"""

import argparse
import asyncio
import hashlib
import io
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from tempfile import gettempdir

from sqlalchemy import delete, func, select
from sqlalchemy.engine import make_url

from tests.environment import validate_loaded_test_settings
from tests.environment import validate_test_environment as validate_base_environment


FIXTURE_ID = "CODEX-CODE-WOPI-20261011"
USERNAME = "CODEX-office-20261011"
ORIGINAL_MARKER = "CODEX-ORIGINAL-20261011"
SAVED_MARKER = "CODEX-SAVED-20261011"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_FILE_BYTES = 20 * 1024 * 1024
DATABASE_NAME = re.compile(r"oa_test_[A-Za-z0-9_]+\.(?:db|sqlite|sqlite3)\Z")
SAFE_COMPONENT = re.compile(r"[A-Za-z0-9_-]+\Z")


@dataclass(frozen=True)
class TestEnvironment:
    database_path: Path
    upload_root: Path
    volume_root: Path


def test_volume_path(raw_path: str, label: str) -> tuple[Path, Path]:
    path = Path(raw_path).expanduser()
    if not path.is_absolute():
        raise RuntimeError(f"{label} 必须是隔离卷内的绝对路径")
    path = path.resolve()
    runtime_root = Path(__file__).resolve().parent / ".runtime"
    if runtime_root.resolve() != runtime_root:
        raise RuntimeError("tests/.runtime 不能指向测试目录以外")
    if path != runtime_root and path.is_relative_to(runtime_root):
        return path, runtime_root
    temp_root = Path(gettempdir()).resolve()
    if not path.is_relative_to(temp_root):
        raise RuntimeError(f"{label} 必须位于当前 tests/.runtime 或系统临时目录的 oa-test-* 子目录")
    parts = path.relative_to(temp_root).parts
    if len(parts) < 2 or not parts[0].startswith("oa-test-") or len(parts[0]) <= len("oa-test-"):
        raise RuntimeError(f"{label} 必须位于 oa-test-* 专属卷内部")
    return path, temp_root / parts[0]


def validate_test_environment() -> TestEnvironment:
    """在导入任何应用配置或连接数据库前执行更严格的隔离门禁。"""
    values = validate_base_environment()
    if values["APP_ENV"].strip().lower() not in {"test", "testing"}:
        raise RuntimeError("本 fixture 只允许 APP_ENV=test/testing")
    url = make_url(values["DATABASE_URL"])
    if url.drivername != "sqlite+aiosqlite" or url.database == ":memory:" or url.username or url.password:
        raise RuntimeError("本 fixture 只允许独立、文件型 SQLite 测试库")
    database, volume = test_volume_path(url.database or "", "DATABASE_URL")
    uploads, upload_volume = test_volume_path(values["UPLOAD_ROOT"], "UPLOAD_ROOT")
    if not DATABASE_NAME.fullmatch(database.name) or upload_volume != volume or database.is_relative_to(uploads):
        raise RuntimeError("数据库必须命名 oa_test_*.db/sqlite/sqlite3，与上传目录同卷且不在上传目录内")
    if Path(".env").exists():
        raise RuntimeError("fixture 不读取 .env；请在无 .env 的隔离测试入口显式传入环境变量")
    for key in os.environ:
        upper = key.upper()
        cloud_key = upper.startswith(("OSS_", "ALIYUN_", "ALIBABA_CLOUD_", "AWS_", "MINIO_"))
        credential_key = any(word in upper for word in ("CREDENTIAL", "ACCESS_KEY", "SECRET", "TOKEN", "PASSWORD"))
        if cloud_key and credential_key and os.environ[key].strip():
            raise RuntimeError("fixture 拒绝携带 OSS 或对象存储凭据的环境")
    return TestEnvironment(database, uploads, volume)


def guarded_file(raw_path: str, root: Path) -> Path:
    path = Path(raw_path)
    if not path.is_absolute():
        raise RuntimeError("fixture 文件必须使用绝对路径")
    resolved = path.resolve()
    if not resolved.is_relative_to(root) or resolved == root:
        raise RuntimeError("fixture 文件不在授权的隔离目录内")
    if path.is_symlink() or any(parent.is_symlink() for parent in path.parents if parent != root.parent):
        raise RuntimeError("fixture 不接受符号链接文件或目录")
    return resolved


def metadata_path(raw_path: str, environment: TestEnvironment) -> Path:
    path = guarded_file(raw_path, environment.volume_root)
    if path == environment.database_path or path.is_relative_to(environment.upload_root) or path.suffix != ".json":
        raise RuntimeError("metadata 必须是测试卷内、数据库及上传目录之外的 JSON 文件")
    return path


def load_metadata(path: Path, environment: TestEnvironment) -> dict:
    if path.stat().st_size > 16 * 1024:
        raise RuntimeError("fixture metadata 超过大小限制")
    payload = json.loads(path.read_text(encoding="utf-8"))
    expected_keys = {"fixture_id", "username", "case_id", "attachment_id", "user_id", "database_path", "upload_root", "original_path", "original_sha256"}
    if not isinstance(payload, dict) or set(payload) != expected_keys:
        raise RuntimeError("fixture metadata 字段不符合专属合同")
    if payload["fixture_id"] != FIXTURE_ID or payload["username"] != USERNAME:
        raise RuntimeError("metadata 不是本任务的 fixture")
    if any(type(payload[key]) is not int or payload[key] <= 0 for key in ("case_id", "attachment_id", "user_id")):
        raise RuntimeError("metadata 对象 ID 无效")
    if payload["database_path"] != str(environment.database_path) or payload["upload_root"] != str(environment.upload_root):
        raise RuntimeError("metadata 与当前隔离数据库或上传目录不一致")
    original = guarded_file(payload["original_path"], environment.upload_root)
    if original != environment.upload_root / f"{FIXTURE_ID}-original.docx":
        raise RuntimeError("metadata 原件路径不属于本 fixture")
    if not re.fullmatch(r"[0-9a-f]{64}", payload["original_sha256"]):
        raise RuntimeError("metadata 原件摘要无效")
    return payload


async def load_owned_rows(db, metadata: dict, *, allow_missing: bool = False):
    from app.models import BusinessRecord, FileAttachment, User

    case = await db.get(BusinessRecord, metadata["case_id"])
    item = await db.get(FileAttachment, metadata["attachment_id"])
    user = await db.get(User, metadata["user_id"])
    if not allow_missing and any(row is None for row in (case, item, user)):
        raise RuntimeError("fixture 对象缺失，不能生成验收回执")
    if case is not None and (case.module != "case" or case.serial_no != FIXTURE_ID or case.owner != USERNAME or (case.data or {}).get("code_wopi_browser_fixture") != FIXTURE_ID):
        raise RuntimeError("案件 ID 已不属于本 fixture，拒绝操作")
    if item is not None and (item.record_id != metadata["case_id"] or item.uploader != USERNAME or item.original_name != f"{FIXTURE_ID}.docx"):
        raise RuntimeError("附件 ID 已不属于本 fixture，拒绝操作")
    if user is not None and (user.username != USERNAME or (user.profile or {}).get("code_wopi_browser_fixture") != FIXTURE_ID):
        raise RuntimeError("账号 ID 已不属于本 fixture，拒绝操作")
    return case, item, user


def session_entries(case, attachment_id: int) -> list[dict]:
    sessions = (case.data or {}).get("libreoffice_sessions", {}) if case else {}
    if not isinstance(sessions, dict):
        raise RuntimeError("fixture 会话字段不是预期 JSON 对象")
    entries = []
    for session_id, session in sessions.items():
        if not isinstance(session, dict) or session.get("attachment_id") != attachment_id or session.get("user") != USERNAME or session.get("id") != session_id or not SAFE_COMPONENT.fullmatch(session_id):
            raise RuntimeError("fixture 出现非本任务会话，拒绝清理")
        entries.append(session)
    return entries


def fixture_files(environment: TestEnvironment, metadata: dict, case, item) -> set[Path]:
    original = guarded_file(metadata["original_path"], environment.upload_root)
    revision_root = guarded_file(str(environment.upload_root / "libreoffice" / str(metadata["attachment_id"])), environment.upload_root)
    paths = {original}
    references = [item.path if item else None]
    for entry in session_entries(case, metadata["attachment_id"]):
        references.extend([entry.get("path"), (entry.get("original_storage") or {}).get("path")])
        if entry.get("pending"):
            references.extend([entry["pending"].get("path"), entry["pending"].get("base_path")])
    documents = (case.data or {}).get("libreoffice_documents", {}) if case else {}
    if not isinstance(documents, dict) or set(documents) - {str(metadata["attachment_id"])}:
        raise RuntimeError("fixture 出现其他附件的版本历史，拒绝扩大清理范围")
    for document in documents.values():
        references.extend([document["path"], document["original_storage"]["path"]])
        for version in document["versions"]:
            references.extend([version["path"], version["previous_path"]])
    for raw_path in references:
        if not raw_path:
            continue
        path = guarded_file(raw_path, environment.upload_root)
        if path != original and not path.is_relative_to(revision_root):
            raise RuntimeError("附件或会话引用了本 fixture 范围之外的文件")
        paths.add(path)
    # 仅枚举本附件的专属版本目录，也包含写文件成功但事务失败的孤立版本。
    if revision_root.exists():
        for folder in revision_root.iterdir():
            guarded_file(str(folder), environment.upload_root)
            if not folder.is_dir() or not SAFE_COMPONENT.fullmatch(folder.name):
                raise RuntimeError("本附件版本目录出现非预期条目")
            for path in folder.iterdir():
                guarded_file(str(path), environment.upload_root)
                if not path.is_file() or path.suffix != ".docx" or not SAFE_COMPONENT.fullmatch(path.stem):
                    raise RuntimeError("本附件版本目录出现非 DOCX 文件")
                paths.add(path.resolve())
    return paths


async def setup(db, environment: TestEnvironment, path: Path, password: str) -> dict:
    from docx import Document
    from app.models import BusinessRecord, FileAttachment, User
    from app.security import hash_password

    if not password or len(password) < 12:
        raise RuntimeError("setup 要求主控提供至少 12 字符的非生产 dummy 密码")
    if path.exists() or await db.scalar(select(BusinessRecord.id).where(BusinessRecord.serial_no == FIXTURE_ID)) or await db.scalar(select(User.id).where(User.username == USERNAME)):
        raise RuntimeError("fixture 已存在；禁止覆盖，请先用原 metadata 精确清理")
    original = environment.upload_root / f"{FIXTURE_ID}-original.docx"
    guarded_file(str(original), environment.upload_root)
    document = Document()
    document.add_paragraph(ORIGINAL_MARKER)
    document.add_paragraph("中文原文：隔离验收文档，保留原件并验证真实在线保存。")
    document.add_paragraph("English original paragraph: isolated CODE edit-save-reopen acceptance.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "项目 / Item", "原始内容 / Original"
    table.cell(1, 0).text, table.cell(1, 1).text = "表格验收 / Table", ORIGINAL_MARKER
    content = io.BytesIO()
    document.save(content)
    raw = content.getvalue()
    created_paths = []
    try:
        user = User(username=USERNAME, display_name="CODEX Office 隔离验收", department=FIXTURE_ID,
                    password_hash=hash_password(password), role="admin", role_ids=["admin"], is_active=True,
                    must_change_password=False, profile={"code_wopi_browser_fixture": FIXTURE_ID})
        case = BusinessRecord(module="case", serial_no=FIXTURE_ID, title="CODEX CODE WOPI 隔离验收",
                              customer="", status="文书准备", owner=USERNAME, department=FIXTURE_ID,
                              description="仅用于本任务隔离浏览器验收",
                              data={"code_wopi_browser_fixture": FIXTURE_ID, "case_type": "民事案件", "case_creation_step": "completed"})
        db.add_all([user, case])
        await db.flush()
        item = FileAttachment(record_id=case.id, category="法院诉讼文书", original_name=f"{FIXTURE_ID}.docx",
                              stored_name=original.name, content_type=CONTENT_TYPE, size=len(raw),
                              path=str(original), uploader=USERNAME, remark=FIXTURE_ID)
        db.add(item)
        await db.flush()
        if (environment.upload_root / "libreoffice" / str(item.id)).exists():
            raise RuntimeError("新附件 ID 存在历史版本目录，拒绝复用")
        metadata = {"fixture_id": FIXTURE_ID, "username": USERNAME, "case_id": case.id,
                    "attachment_id": item.id, "user_id": user.id, "database_path": str(environment.database_path),
                    "upload_root": str(environment.upload_root), "original_path": str(original),
                    "original_sha256": hashlib.sha256(raw).hexdigest()}
        environment.upload_root.mkdir(parents=True, exist_ok=True)
        path.parent.mkdir(parents=True, exist_ok=True)
        with original.open("xb") as handle:
            created_paths.append(original)
            handle.write(raw)
        with path.open("x", encoding="utf-8") as handle:
            created_paths.append(path)
            json.dump(metadata, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        await db.commit()
    except Exception:
        await db.rollback()
        for created in reversed(created_paths):
            created.unlink(missing_ok=True)
        raise
    return {"operation": "setup", **metadata, "metadata_path": str(path), "storage_backend": "local", "route": f"/?page=case-detail-{case.id}-{FIXTURE_ID}"}


def read_file(path: Path) -> bytes:
    with path.open("rb") as handle:
        raw = handle.read(MAX_FILE_BYTES + 1)
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise RuntimeError("fixture DOCX 为空或超过大小限制")
    return raw


async def counts(db, metadata: dict) -> dict:
    from app.models import BusinessRecord, CaseEvent, FileAttachment, User, WorkflowEvent

    filters = {"records": (BusinessRecord, BusinessRecord.serial_no == FIXTURE_ID),
               "attachments": (FileAttachment, FileAttachment.record_id == metadata["case_id"]),
               "workflow_events": (WorkflowEvent, WorkflowEvent.record_id == metadata["case_id"]),
               "case_events": (CaseEvent, CaseEvent.case_record_id == metadata["case_id"]),
               "test_accounts": (User, User.username == USERNAME)}
    result = {}
    for key, (model, predicate) in filters.items():
        result[key] = await db.scalar(select(func.count()).select_from(model).where(predicate))
    return result


async def report(db, environment: TestEnvironment, metadata: dict, saved_marker: str) -> dict:
    from docx import Document

    case, item, _user = await load_owned_rows(db, metadata)
    paths = fixture_files(environment, metadata, case, item)
    original = read_file(Path(metadata["original_path"]))
    current = read_file(guarded_file(item.path, environment.upload_root))
    document = Document(io.BytesIO(current))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    texts = paragraphs + [text for table in tables for row in table for text in row]
    sessions = session_entries(case, item.id)
    expiry = item.word_editor_lock_expires_at
    if expiry is not None and expiry.tzinfo is None:
        expiry = expiry.replace(tzinfo=timezone.utc)
    return {"operation": "report", "fixture_id": FIXTURE_ID, "case_id": case.id, "attachment_id": item.id,
            "current_path": item.path, "paragraphs": paragraphs, "tables": tables,
            "saved_marker": saved_marker, "saved_marker_found": any(saved_marker in text for text in texts),
            "current_sha256": hashlib.sha256(current).hexdigest(), "original_sha256": hashlib.sha256(original).hexdigest(),
            "original_hash_unchanged": hashlib.sha256(original).hexdigest() == metadata["original_sha256"],
            "current_differs_from_original": current != original, "session_count": len(sessions),
            "session_statuses": [entry.get("status") for entry in sessions],
            "sessions": [{"status": entry.get("status"), "save_status": entry.get("save_status"),
                          "error": entry.get("error"), "last_saved_at": entry.get("last_saved_at"),
                          "wopi_lock": bool(entry.get("wopi_lock"))} for entry in sessions],
            "lock_count": int(bool(item.word_editor_lock_token)),
            "active_lock_count": int(bool(item.word_editor_lock_token) and expiry is not None and expiry > datetime.now(timezone.utc)),
            "version_file_count": sum(path.is_file() for path in paths if str(path) != metadata["original_path"]),
            "database_counts": await counts(db, metadata)}


async def cleanup(db, environment: TestEnvironment, metadata: dict, path: Path) -> dict:
    from app.models import BusinessRecord, CaseEvent, FileAttachment, User, WorkflowEvent

    case, item, _user = await load_owned_rows(db, metadata, allow_missing=True)
    paths = fixture_files(environment, metadata, case, item)
    sessions = session_entries(case, metadata["attachment_id"])
    expiry = item.word_editor_lock_expires_at if item else None
    if expiry is not None and expiry.tzinfo is None:
        expiry = expiry.replace(tzinfo=timezone.utc)
    if item and item.word_editor_lock_token and expiry and expiry > datetime.now(timezone.utc):
        raise RuntimeError("fixture 仍持有有效 shared Word 租约；须先真实关闭编辑器，不能由 cleanup 强制解锁")
    if any(entry.get("status") != "closed" and datetime.fromisoformat(entry["expires_at"].replace("Z", "+00:00")) > datetime.now(timezone.utc) for entry in sessions):
        raise RuntimeError("请先关闭本 fixture 编辑器并确认会话 closed/expired，再执行清理")
    if await db.scalar(select(FileAttachment.id).where(FileAttachment.record_id == metadata["case_id"], FileAttachment.id != metadata["attachment_id"])):
        raise RuntimeError("fixture 案件存在额外附件，拒绝扩大清理范围")
    if await db.scalar(select(BusinessRecord.id).where(BusinessRecord.owner == USERNAME, BusinessRecord.id != metadata["case_id"])) or await db.scalar(select(FileAttachment.id).where(FileAttachment.uploader == USERNAME, FileAttachment.id != metadata["attachment_id"])):
        raise RuntimeError("fixture 账号关联了额外业务数据，拒绝扩大清理范围")
    if paths and await db.scalar(select(FileAttachment.id).where(FileAttachment.path.in_([str(file) for file in paths]), FileAttachment.id != metadata["attachment_id"])):
        raise RuntimeError("fixture 文件被其他附件引用，拒绝删除共享文件")
    # 仅清理通过 ID 与任务标识双重核验的对象；原默认 admin 不参与任何写入。
    await db.execute(delete(WorkflowEvent).where(WorkflowEvent.record_id == metadata["case_id"]))
    await db.execute(delete(CaseEvent).where(CaseEvent.case_record_id == metadata["case_id"]))
    await db.execute(delete(FileAttachment).where(FileAttachment.id == metadata["attachment_id"]))
    await db.execute(delete(BusinessRecord).where(BusinessRecord.id == metadata["case_id"]))
    await db.execute(delete(User).where(User.id == metadata["user_id"]))
    await db.commit()
    for file in paths:
        file.unlink(missing_ok=True)
    revision_root = environment.upload_root / "libreoffice" / str(metadata["attachment_id"])
    if revision_root.exists():
        for folder in revision_root.iterdir():
            folder.rmdir()
        revision_root.rmdir()
    residual = await counts(db, metadata)
    remaining_case = await db.get(BusinessRecord, metadata["case_id"])
    residual.update({"sessions": len(session_entries(remaining_case, metadata["attachment_id"])),
                     "document_histories": len((remaining_case.data or {}).get("libreoffice_documents", {})) if remaining_case else 0,
                     "locks": int(bool((await db.get(FileAttachment, metadata["attachment_id"])))),
                     "files": sum(file.exists() for file in paths)})
    if any(residual.values()):
        raise RuntimeError("fixture 清理回查仍有残留；保留 metadata 以继续精确排查")
    path.unlink()
    return {"operation": "cleanup", "fixture_id": FIXTURE_ID, "zero_residual": True,
            "residual": residual, "metadata_removed": not path.exists(), "database_file_preserved": True}


async def run(args) -> dict:
    environment = validate_test_environment()
    path = metadata_path(args.metadata, environment)
    if args.command != "setup" and not environment.database_path.is_file():
        raise RuntimeError("隔离测试数据库不存在，拒绝创建空库掩盖缺失")
    from app.config import settings

    validate_loaded_test_settings(settings)
    if settings.oss_credentials_file or settings.minio_access_key or settings.minio_secret_key:
        raise RuntimeError("已加载设置含对象存储凭据，拒绝连接")
    from app.database import Base, SessionLocal, engine
    import app.models  # 注册现有完整模型，不导入 app.main 或初始化默认账号。

    try:
        if args.command == "setup":
            environment.database_path.parent.mkdir(parents=True, exist_ok=True)
            async with engine.begin() as connection:
                await connection.run_sync(Base.metadata.create_all)
        async with SessionLocal() as db:
            if args.command == "setup":
                return await setup(db, environment, path, args.password)
            metadata = load_metadata(path, environment)
            if args.command == "report":
                return await report(db, environment, metadata, args.saved_marker)
            return await cleanup(db, environment, metadata, path)
    finally:
        await engine.dispose()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    for command in ("setup", "report", "cleanup"):
        child = commands.add_parser(command)
        child.add_argument("--metadata", required=True)
        if command == "setup":
            child.add_argument("--password", required=True)
        if command == "report":
            child.add_argument("--saved-marker", default=SAVED_MARKER)
    args = parser.parse_args()
    try:
        result = asyncio.run(run(args))
    except (RuntimeError, ValueError, OSError, KeyError) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    except Exception:
        # 数据库异常可能带 SQL 参数，不能将密码哈希或会话 token 输出到日志。
        print(json.dumps({"error": "fixture 执行失败；未输出数据库参数或敏感配置"}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
