"""领域路由拆分后，导入查询和附件文件事务的双数据库契约。"""

import asyncio
import csv
import errno
import hashlib
import io
import os
import tempfile
import threading
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from fastapi import HTTPException
from httpx import ASGITransport, AsyncClient
from docx import Document
from sqlalchemy import event, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.areas.system import agent_documents, attachments
from app.areas.legal import seals
from app.config import settings
from app.core import attachment_deletion, formatters, storage
from app.core.dependencies import current_identity, get_db
from app.database import Base
from app.main import app
from app.models import AgentDocument, BusinessRecord, DocumentTemplate, FileAttachment, User, WorkflowEvent
from app.models_shared import AttachmentBatchInput


POSTGRES_URL = os.environ.get("AGENT_COMMANDS_TEST_POSTGRES_URL", "")


class DomainRouteStorageAssertions:
    async def asyncSetUp(self):
        if self.database_url:
            if make_url(self.database_url).database != "oa_test_enterprise_commands":
                raise RuntimeError("领域路由测试只能使用专用 PostgreSQL 测试库")
            self.schema = f"oa_test_enterprise_domain_{uuid4().hex[:12]}"
            self.admin_engine = create_async_engine(self.database_url)
            async with self.admin_engine.begin() as connection:
                await connection.execute(text(f"CREATE SCHEMA {self.schema}"))
            self.engine = create_async_engine(
                self.database_url,
                connect_args={"server_settings": {"search_path": self.schema}},
            )
        else:
            self.engine = create_async_engine(
                "sqlite+aiosqlite:///:memory:",
                connect_args={"check_same_thread": False},
                poolclass=StaticPool,
            )
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.temporary = tempfile.TemporaryDirectory(prefix="oa-test-enterprise-domain-")
        self.upload_root = Path(self.temporary.name)
        self.patchers = [
            patch.object(attachment_deletion, "UPLOAD_ROOT", self.upload_root),
            patch.object(attachments, "UPLOAD_ROOT", self.upload_root),
            patch.object(storage, "UPLOAD_ROOT", self.upload_root),
            patch.object(formatters, "UPLOAD_ROOT", self.upload_root),
            patch.object(agent_documents, "UPLOAD_ROOT", self.upload_root),
            patch.object(seals, "UPLOAD_ROOT", self.upload_root),
        ]
        for patcher in self.patchers:
            patcher.start()
        self.identity = {
            "username": "domain-admin", "role": "admin", "_actual_role_ids": ["admin"],
            "role_ids": ["admin"], "menu_keys": [], "action_keys": [],
            "_page_menu_capability": False,
        }

        async def test_db():
            async with self.sessions() as db:
                yield db

        app.dependency_overrides[current_identity] = lambda: self.identity
        app.dependency_overrides[get_db] = test_db
        self.client = AsyncClient(transport=ASGITransport(app=app), base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()
        app.dependency_overrides.clear()
        for patcher in reversed(self.patchers):
            patcher.stop()
        self.temporary.cleanup()
        await self.engine.dispose()
        if self.database_url:
            async with self.admin_engine.begin() as connection:
                await connection.execute(text(f"DROP SCHEMA {self.schema} CASCADE"))
            await self.admin_engine.dispose()

    async def test_record_import_keeps_partial_results_and_only_loads_references(self):
        async with self.sessions() as db:
            db.add(User(username="domain-admin", display_name="领域管理员", department="上海", password_hash="x", role="admin", role_ids=["admin"]))
            customer = BusinessRecord(
                module="customer", serial_no="ENTERPRISE-DOMAIN-CUSTOMER", title="客户甲",
                customer="", status="正常", owner="domain-admin", department="上海", data={},
            )
            db.add(customer)
            await db.flush()
            case = BusinessRecord(
                module="case", serial_no="ENTERPRISE-DOMAIN-CASE", title="案件甲",
                customer="客户甲", status="办理中", owner="domain-admin", department="上海",
                data={"customer_id": customer.id, "customer_no": customer.serial_no, "customer_title": customer.title},
            )
            db.add(case)
            db.add_all(BusinessRecord(
                module="case", serial_no=f"ENTERPRISE-DOMAIN-UNRELATED-{index:03d}",
                title=f"无关案件{index}", customer="无关客户", status="办理中",
                owner="domain-admin", department="上海", data={"large": "x" * 4096},
            ) for index in range(120))
            await db.commit()

        payload = io.StringIO()
        writer = csv.writer(payload)
        writer.writerow(["业务编号", "标题", "收发类型", "关联案号", "文件日期"])
        writer.writerow(["ENTERPRISE-DOMAIN-DOC-1", "文书甲", "收文", "ENTERPRISE-DOMAIN-CASE", "2026-09-30"])
        writer.writerow(["ENTERPRISE-DOMAIN-DOC-2", "文书乙", "收文", "MISSING-CASE", "2026-09-30"])
        writer.writerow(["ENTERPRISE-DOMAIN-CASE", "文书丙", "收文", "ENTERPRISE-DOMAIN-CASE", "2026-09-30"])
        loaded = []

        def on_load(target, context):
            loaded.append(target.id)

        event.listen(BusinessRecord, "load", on_load)
        try:
            response = await self.client.post(
                "/api/v1/records/import", params={"module": "document"},
                files={"file": ("records.csv", payload.getvalue().encode("utf-8"), "text/csv")},
            )
        finally:
            event.remove(BusinessRecord, "load", on_load)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual((response.json()["created"], response.json()["failed"]), (1, 2))
        self.assertLess(len(loaded), 10, "不应把无关业务记录及大 JSON 加载为 ORM 实体")
        async with self.sessions() as db:
            document = await db.scalar(select(BusinessRecord).where(BusinessRecord.serial_no == "ENTERPRISE-DOMAIN-DOC-1"))
            self.assertIsNotNone(document)
            self.assertEqual(document.data["case_id"], case.id)
            self.assertEqual(document.data["customer_id"], customer.id)
            self.assertEqual(document.customer, "客户甲")
            self.assertEqual(len((await db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == document.id))).all()), 1)
            self.assertIsNone(await db.scalar(select(BusinessRecord).where(BusinessRecord.serial_no == "ENTERPRISE-DOMAIN-DOC-2")))
        exported = await self.client.get("/api/v1/records/export", params={"module": "document"})
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertIn("ENTERPRISE-DOMAIN-DOC-1", exported.text)
        listed = await self.client.get("/api/v1/records", params={"module": "document"})
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(listed.json()["total"], 1)
        self.assertEqual(listed.json()["items"][0]["serial_no"], "ENTERPRISE-DOMAIN-DOC-1")

    async def test_attachment_delete_success_failure_and_reconciliation(self):
        path = self.upload_root / f"enterprise-domain-{uuid4().hex}.txt"
        path.write_bytes(b"attachment body")
        async with self.sessions() as db:
            item = FileAttachment(
                original_name="证据.txt", stored_name=path.name, content_type="text/plain",
                size=path.stat().st_size, path=str(path), uploader="domain-admin",
            )
            db.add(item)
            await db.commit()
            attachment_id = item.id
        response = await self.client.delete(f"/api/v1/attachments/{attachment_id}")
        self.assertEqual(response.status_code, 204, response.text)
        self.assertFalse(path.exists())
        self.assertEqual(list(self.upload_root.glob(".pending-delete-*")), [])
        async with self.sessions() as db:
            self.assertIsNone(await db.get(FileAttachment, attachment_id))

        shared_path = self.upload_root / f"enterprise-shared-{uuid4().hex}.txt"
        shared_path.write_bytes(b"shared attachment body")
        async with self.sessions() as db:
            first = FileAttachment(original_name="共享甲.txt", stored_name=shared_path.name, content_type="text/plain", size=shared_path.stat().st_size, path=str(shared_path), uploader="domain-admin")
            second = FileAttachment(original_name="共享乙.txt", stored_name=f"shared-alias-{uuid4().hex}.txt", content_type="text/plain", size=shared_path.stat().st_size, path=str(shared_path), uploader="domain-admin")
            db.add_all([first, second])
            await db.commit()
            first_id, second_id = first.id, second.id
        self.assertEqual((await self.client.delete(f"/api/v1/attachments/{first_id}")).status_code, 204)
        self.assertEqual(shared_path.read_bytes(), b"shared attachment body")
        async with self.sessions() as db:
            self.assertIsNone(await db.get(FileAttachment, first_id))
            self.assertIsNotNone(await db.get(FileAttachment, second_id))
        self.assertEqual((await self.client.delete(f"/api/v1/attachments/{second_id}")).status_code, 204)
        self.assertFalse(shared_path.exists())

        path.write_bytes(b"rollback body")
        async with self.sessions() as db:
            item = FileAttachment(
                original_name="回滚.txt", stored_name=path.name, content_type="text/plain",
                size=path.stat().st_size, path=str(path), uploader="domain-admin",
            )
            db.add(item)
            await db.commit()
            rollback_id = item.id
        async with self.sessions() as db:
            with patch.object(db, "commit", side_effect=RuntimeError("injected commit failure")):
                with self.assertRaisesRegex(RuntimeError, "injected commit failure"):
                    await attachments.delete_attachment(rollback_id, self.identity, db)
        self.assertEqual(path.read_bytes(), b"rollback body")
        async with self.sessions() as db:
            self.assertIsNotNone(await db.get(FileAttachment, rollback_id))

        original_unlink = Path.unlink

        def fail_staged_unlink(candidate, *args, **kwargs):
            if candidate.name.startswith(".pending-delete-") and not candidate.name.startswith(".pending-delete-lock-"):
                raise OSError("injected unlink failure")
            return original_unlink(candidate, *args, **kwargs)

        async with self.sessions() as db:
            with patch.object(Path, "unlink", fail_staged_unlink):
                with self.assertRaises(HTTPException) as failure:
                    await attachments.delete_attachment(rollback_id, self.identity, db)
        self.assertEqual(failure.exception.status_code, 500)
        staged = [item for item in self.upload_root.glob(".pending-delete-*") if not item.name.startswith(".pending-delete-lock-")]
        self.assertEqual(len(staged), 1)
        pending_locks = list(self.upload_root.glob(".pending-delete-lock-*"))
        self.assertEqual(len(pending_locks), 1, "暂存正文仍存在时必须保留同一把协调锁")
        self.assertFalse(path.exists())
        async with self.sessions() as db:
            self.assertIsNone(await db.get(FileAttachment, rollback_id))
            with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertFalse(staged[0].exists())
        self.assertFalse(pending_locks[0].exists())

        path.write_bytes(b"orphan lock body")
        async with self.sessions() as db:
            item = FileAttachment(
                original_name="孤立锁.txt", stored_name=path.name, content_type="text/plain",
                size=path.stat().st_size, path=str(path), uploader="domain-admin",
            )
            db.add(item)
            await db.commit()
            orphan_id = item.id

        def fail_lock_unlink(candidate, *args, **kwargs):
            if candidate.name.startswith(".pending-delete-lock-"):
                raise OSError("injected lock unlink failure")
            return original_unlink(candidate, *args, **kwargs)

        async with self.sessions() as db:
            with patch.object(Path, "unlink", fail_lock_unlink):
                with self.assertRaises(HTTPException) as failure:
                    await attachments.delete_attachment(orphan_id, self.identity, db)
        self.assertEqual(failure.exception.status_code, 500)
        self.assertFalse(path.exists())
        orphan_locks = list(self.upload_root.glob(".pending-delete-lock-*"))
        self.assertEqual(len(orphan_locks), 1)
        async with self.sessions() as db:
            self.assertIsNone(await db.get(FileAttachment, orphan_id))
            with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertFalse(orphan_locks[0].exists())

        path.write_bytes(b"crash before commit")
        async with self.sessions() as db:
            item = FileAttachment(
                original_name="恢复.txt", stored_name=path.name, content_type="text/plain",
                size=path.stat().st_size, path=str(path), uploader="domain-admin",
            )
            db.add(item)
            await db.commit()
            recover_id = item.id
        staged = attachment_deletion.stage_attachment_delete(recover_id, path)
        self.assertIsNotNone(staged)
        unrelated = self.upload_root / ".pending-delete-not-a-valid-record"
        unrelated.write_bytes(b"unrelated")
        with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
            async with self.sessions() as db:
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertTrue(staged.staged.exists(), "另一个实例持有文件锁时不得恢复正在提交的附件")
        self.assertEqual(unrelated.read_bytes(), b"unrelated")
        attachment_deletion._unlock_file(staged.lock_file)
        async with self.sessions() as db:
            with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertEqual(path.read_bytes(), b"crash before commit")

        shared_path.write_bytes(b"crash with shared reference")
        async with self.sessions() as db:
            first = FileAttachment(original_name="崩溃甲.txt", stored_name=shared_path.name, content_type="text/plain", size=shared_path.stat().st_size, path=str(shared_path), uploader="domain-admin")
            second = FileAttachment(original_name="崩溃乙.txt", stored_name=f"shared-alias-{uuid4().hex}.txt", content_type="text/plain", size=shared_path.stat().st_size, path=str(shared_path), uploader="domain-admin")
            db.add_all([first, second])
            await db.commit()
            first_id, second_id = first.id, second.id
        staged = attachment_deletion.stage_attachment_delete(first_id, shared_path)
        self.assertIsNotNone(staged)
        async with self.sessions() as db:
            await db.delete(await db.get(FileAttachment, first_id))
            await db.commit()
        attachment_deletion._unlock_file(staged.lock_file)
        with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
            async with self.sessions() as db:
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertEqual(shared_path.read_bytes(), b"crash with shared reference")
        async with self.sessions() as db:
            self.assertIsNotNone(await db.get(FileAttachment, second_id))

    async def test_attachment_lock_faults_are_not_reported_as_contention(self):
        lock_path = self.upload_root / ".pending-delete-lock-1-0123456789abcdef0123456789abcdef-lock.txt"
        with patch.object(attachment_deletion.os, "open", side_effect=OSError(errno.EIO, "injected open failure")):
            with self.assertRaisesRegex(OSError, "injected open failure"):
                attachment_deletion._lock_file(lock_path)

        original_fdopen = attachment_deletion.os.fdopen

        class FailingWriteHandle:
            def __init__(self, wrapped):
                self.wrapped = wrapped

            def __getattr__(self, name):
                return getattr(self.wrapped, name)

            def write(self, value):
                raise OSError(errno.EACCES, "injected write failure")

        with patch.object(attachment_deletion.os, "fdopen", side_effect=lambda fd, mode: FailingWriteHandle(original_fdopen(fd, mode))):
            with self.assertRaisesRegex(OSError, "injected write failure"):
                attachment_deletion._lock_file(lock_path)
        handle = attachment_deletion._lock_file(lock_path)
        self.assertIsNotNone(handle, "写入失败后必须释放锁")
        attachment_deletion._unlock_file(handle)
        lock_path.unlink()
        outside_target = self.upload_root / "outside-lock-target.txt"
        outside_target.write_bytes(b"untouched")
        lock_path.symlink_to(outside_target)
        with self.assertRaisesRegex(OSError, "符号链接"):
            attachment_deletion._lock_file(lock_path)
        self.assertEqual(outside_target.read_bytes(), b"untouched")

    async def test_lock_extension_upload_crash_restores_original_file(self):
        async with self.sessions() as db:
            db.add(User(username="domain-admin", display_name="领域管理员", department="上海", password_hash="x", role="admin", role_ids=["admin"]))
            case = BusinessRecord(module="case", serial_no="ENTERPRISE-DOMAIN-LOCK-CASE", title="扩展名锁文件案件", customer="客户甲", status="办理中", owner="domain-admin", department="上海", data={})
            db.add(case)
            await db.commit()
            case_id = case.id
        uploaded = await self.client.post(
            "/api/v1/attachments", data={"record_id": str(case_id), "category": "普通附件"},
            files={"file": ("evidence.lock", b"real lock extension payload", "application/octet-stream")},
        )
        self.assertEqual(uploaded.status_code, 201, uploaded.text)
        attachment_id = uploaded.json()["id"]
        async with self.sessions() as db:
            attachment = await db.get(FileAttachment, attachment_id)
            path = Path(attachment.path)
        self.assertEqual(path.suffix, ".lock")
        self.assertEqual(path.read_bytes(), b"real lock extension payload")
        staged = attachment_deletion.stage_attachment_delete(attachment_id, path)
        self.assertIsNotNone(staged)
        attachment_deletion._unlock_file(staged.lock_file)
        with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
            async with self.sessions() as db:
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertEqual(path.read_bytes(), b"real lock extension payload")
        self.assertFalse(staged.staged.exists())
        self.assertFalse(staged.lock_path.exists())
        downloaded = await self.client.get(f"/api/v1/attachments/{attachment_id}/download")
        self.assertEqual(downloaded.status_code, 200, downloaded.text)
        self.assertEqual(downloaded.content, b"real lock extension payload")

    async def test_restore_rejects_broken_original_and_staged_symlinks(self):
        original = self.upload_root / f"enterprise-restore-{uuid4().hex}.txt"
        original.write_bytes(b"preserved attachment")
        staged = attachment_deletion.stage_attachment_delete(1, original)
        self.assertIsNotNone(staged)
        original.symlink_to(self.upload_root / "missing-unrelated-target.txt")
        with self.assertRaises(FileExistsError):
            attachment_deletion.restore_attachment_delete(staged)
        self.assertTrue(original.is_symlink())
        self.assertEqual(staged.staged.read_bytes(), b"preserved attachment")
        self.assertTrue(staged.lock_path.exists())
        original.unlink()
        lock_file = attachment_deletion._lock_file(staged.lock_path)
        self.assertIsNotNone(lock_file)
        attachment_deletion.restore_attachment_delete(
            attachment_deletion.StagedAttachmentDelete(original, staged.staged, staged.lock_path, lock_file),
        )
        self.assertEqual(original.read_bytes(), b"preserved attachment")

        staged = attachment_deletion.stage_attachment_delete(2, original)
        self.assertIsNotNone(staged)
        staged.staged.unlink()
        unrelated = self.upload_root / "unrelated-target.txt"
        unrelated.write_bytes(b"unrelated")
        staged.staged.symlink_to(unrelated)
        with self.assertRaisesRegex(OSError, "暂存文件已变化"):
            attachment_deletion.restore_attachment_delete(staged)
        self.assertFalse(original.exists())
        self.assertTrue(staged.staged.is_symlink())
        self.assertEqual(unrelated.read_bytes(), b"unrelated")
        self.assertTrue(staged.lock_path.exists())
        staged.staged.unlink()
        staged.lock_path.unlink()

    async def test_seal_batch_delete_preserves_shared_paths_and_reports_cleanup_failure(self):
        path = self.upload_root / f"enterprise-seal-{uuid4().hex}.pdf"
        path.write_bytes(b"shared seal source")
        async with self.sessions() as db:
            db.add(User(username="domain-admin", display_name="领域管理员", department="上海", password_hash="x", role="admin", role_ids=["admin"]))
            record = BusinessRecord(module="seal", serial_no="ENTERPRISE-DOMAIN-SEAL", title="批量用印", customer="客户甲", status="草稿", owner="domain-admin", department="上海", data={})
            db.add(record)
            await db.flush()
            first = FileAttachment(record_id=record.id, category="用印文件", original_name="共享甲.pdf", stored_name=path.name, content_type="application/pdf", size=path.stat().st_size, path=str(path), uploader="domain-admin")
            second = FileAttachment(record_id=record.id, category="用印文件", original_name="共享乙.pdf", stored_name=f"alias-{uuid4().hex}.pdf", content_type="application/pdf", size=path.stat().st_size, path=str(path), uploader="domain-admin")
            db.add_all([first, second])
            await db.commit()
            first_id, second_id, record_id = first.id, second.id, record.id

        async with self.sessions() as db:
            result = await seals.batch_delete_seal_attachments(AttachmentBatchInput(attachment_ids=[first_id, second_id]), self.identity, db)
        self.assertEqual(result["deleted"], 2)
        self.assertFalse(path.exists(), "同批次删除所有引用时只能暂存并清理原文件一次")
        async with self.sessions() as db:
            self.assertIsNone(await db.get(FileAttachment, first_id))
            self.assertIsNone(await db.get(FileAttachment, second_id))
            events = (await db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == record_id))).all()
            self.assertEqual(len(events), 2)

        first_path = self.upload_root / f"enterprise-seal-{uuid4().hex}.pdf"
        second_path = self.upload_root / f"enterprise-seal-{uuid4().hex}.pdf"
        first_path.write_bytes(b"first seal")
        second_path.write_bytes(b"second seal")
        async with self.sessions() as db:
            first = FileAttachment(record_id=record_id, category="用印文件", original_name="失败甲.pdf", stored_name=first_path.name, content_type="application/pdf", size=first_path.stat().st_size, path=str(first_path), uploader="domain-admin")
            second = FileAttachment(record_id=record_id, category="用印文件", original_name="失败乙.pdf", stored_name=second_path.name, content_type="application/pdf", size=second_path.stat().st_size, path=str(second_path), uploader="domain-admin")
            db.add_all([first, second])
            await db.commit()
            failure_ids = [first.id, second.id]

        original_unlink = Path.unlink

        def fail_one_cleanup(candidate, *args, **kwargs):
            if candidate.name.startswith(".pending-delete-") and not candidate.name.startswith(".pending-delete-lock-") and candidate.name.endswith(first_path.name):
                raise OSError("injected seal cleanup failure")
            return original_unlink(candidate, *args, **kwargs)

        async with self.sessions() as db:
            with patch.object(Path, "unlink", fail_one_cleanup):
                with self.assertRaises(HTTPException) as failure:
                    await seals.batch_delete_seal_attachments(AttachmentBatchInput(attachment_ids=failure_ids), self.identity, db)
        self.assertEqual(failure.exception.status_code, 500)
        self.assertFalse(first_path.exists())
        self.assertFalse(second_path.exists())
        staged = [item for item in self.upload_root.glob(".pending-delete-*") if not item.name.startswith(".pending-delete-lock-")]
        self.assertEqual(len(staged), 1)
        async with self.sessions() as db:
            for item_id in failure_ids:
                self.assertIsNone(await db.get(FileAttachment, item_id))
            with patch.object(settings, "attachment_delete_reconcile_min_age_seconds", 0):
                await attachment_deletion.reconcile_pending_attachment_deletes(db)
        self.assertFalse(staged[0].exists())

        rollback_path = self.upload_root / f"enterprise-seal-{uuid4().hex}.pdf"
        rollback_path.write_bytes(b"rollback failure seal")
        async with self.sessions() as db:
            item = FileAttachment(record_id=record_id, category="用印文件", original_name="回滚故障.pdf", stored_name=rollback_path.name, content_type="application/pdf", size=rollback_path.stat().st_size, path=str(rollback_path), uploader="domain-admin")
            db.add(item)
            await db.commit()
            rollback_id = item.id
        async with self.sessions() as db:
            with patch.object(db, "commit", side_effect=RuntimeError("injected seal commit failure")):
                with patch.object(db, "rollback", side_effect=RuntimeError("injected seal rollback failure")):
                    with self.assertRaises(ExceptionGroup) as failure:
                        await seals.batch_delete_seal_attachments(AttachmentBatchInput(attachment_ids=[rollback_id]), self.identity, db)
        self.assertEqual(
            {str(error) for error in failure.exception.exceptions},
            {"injected seal commit failure", "injected seal rollback failure"},
        )
        self.assertEqual(rollback_path.read_bytes(), b"rollback failure seal")
        self.assertEqual(list(self.upload_root.glob(".pending-delete-*")), [])
        async with self.sessions() as db:
            self.assertIsNotNone(await db.get(FileAttachment, rollback_id))

    async def test_docx_and_text_preview_leave_event_loop_responsive(self):
        docx_path = self.upload_root / f"enterprise-domain-{uuid4().hex}.docx"
        document = Document()
        document.add_paragraph("文档线程预览")
        document.save(docx_path)
        text_path = self.upload_root / f"enterprise-domain-{uuid4().hex}.txt"
        text_path.write_text("文本线程预览", encoding="utf-8")
        async with self.sessions() as db:
            db.add_all([
                FileAttachment(original_name="文档.docx", stored_name=docx_path.name, content_type="application/octet-stream", size=docx_path.stat().st_size, path=str(docx_path), uploader="domain-admin"),
                FileAttachment(original_name="文本.txt", stored_name=text_path.name, content_type="text/plain", size=text_path.stat().st_size, path=str(text_path), uploader="domain-admin"),
            ])
            await db.commit()
            document_id, text_id = (await db.scalars(select(FileAttachment.id).order_by(FileAttachment.id))).all()
        entered = threading.Event()
        release = threading.Event()
        original_reader = attachments._docx_preview_text

        def slow_reader(path):
            entered.set()
            release.wait(2)
            return original_reader(path)

        async with self.sessions() as db:
            with patch.object(attachments, "_docx_preview_text", slow_reader):
                pending = asyncio.create_task(attachments.preview_attachment(document_id, None, self.identity, db))
                self.assertTrue(await asyncio.to_thread(entered.wait, 2))
                await asyncio.sleep(0.01)
                self.assertFalse(pending.done(), "DOCX 解析必须在线程中等待")
                release.set()
                result = await pending
            self.assertIn("文档线程预览", result["text"])
            original_text_reader = Path.read_text
            entered.clear()
            release.clear()

            def slow_text_reader(path, *args, **kwargs):
                if path.name == text_path.name:
                    entered.set()
                    release.wait(2)
                return original_text_reader(path, *args, **kwargs)

            with patch.object(Path, "read_text", slow_text_reader):
                pending = asyncio.create_task(attachments.preview_attachment(text_id, None, self.identity, db))
                self.assertTrue(await asyncio.to_thread(entered.wait, 2))
                await asyncio.sleep(0.01)
                self.assertFalse(pending.done(), "文本整读必须在线程中等待")
                release.set()
                result = await pending
            self.assertIn("文本线程预览", result["text"])

    async def test_agent_document_and_events_share_one_commit(self):
        async with self.sessions() as db:
            db.add(User(username="domain-admin", display_name="领域管理员", department="上海", password_hash="x", role="admin", role_ids=["admin"]))
            template = DocumentTemplate(name="测试模板", category="案件", fields=["事实"])
            case = BusinessRecord(module="case", serial_no="ENTERPRISE-DOMAIN-AGENT", title="智能文书案件", customer="客户甲", status="办理中", owner="domain-admin", department="上海", data={})
            db.add_all([template, case])
            await db.commit()
            template_id, case_id = template.id, case.id
        response = await self.client.post("/api/v1/agent/documents", json={"template_id": template_id, "record_id": case_id, "title": "智能文书", "instruction": "按模板"})
        self.assertEqual(response.status_code, 201, response.text)
        async with self.sessions() as db:
            items = (await db.scalars(select(AgentDocument))).all()
            self.assertEqual(len(items), 1)
            events = (await db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == case_id))).all()
            self.assertEqual({item.action for item in events}, {"智能文档写入 AI 空间", "创建智能文档"})
            attachment = await db.scalar(select(FileAttachment).where(FileAttachment.record_id == case_id))
            self.assertTrue(Path(attachment.path).is_file())

        from app.models_shared import AgentDocumentInput

        async with self.sessions() as db:
            with patch.object(db, "commit", side_effect=RuntimeError("injected event transaction failure")):
                with self.assertRaisesRegex(RuntimeError, "injected event transaction failure"):
                    await agent_documents.create_agent_document(
                        AgentDocumentInput(template_id=template_id, record_id=case_id, title="失败文书", instruction="按模板"),
                        self.identity, db,
                    )
        async with self.sessions() as db:
            self.assertEqual(len((await db.scalars(select(AgentDocument))).all()), 1)
            self.assertEqual(len((await db.scalars(select(FileAttachment).where(FileAttachment.record_id == case_id))).all()), 1)
            self.assertEqual(len((await db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == case_id))).all()), 2)
        self.assertEqual(len(list(self.upload_root.glob("*.docx"))), 1)

    async def test_agent_document_download_and_writeback_cleanup_failures(self):
        async with self.sessions() as db:
            db.add(User(username="domain-admin", display_name="领域管理员", department="上海", password_hash="x", role="admin", role_ids=["admin"]))
            template = DocumentTemplate(name="回写模板", category="案件", fields=["内容"])
            record = BusinessRecord(module="case", serial_no="ENTERPRISE-DOMAIN-WRITEBACK", title="智能文书回写案件", customer="客户甲", status="办理中", owner="domain-admin", department="上海", data={})
            db.add_all([template, record])
            await db.flush()
            document_ids = []
            for index in range(5):
                content = f"正文第 {index} 份\n- 证据材料"
                item = AgentDocument(
                    job_no=f"ENTERPRISE-DOMAIN-DOC-{uuid4().hex}", template_id=template.id,
                    record_id=record.id, title=f"回写文书 {index}", content=content,
                    status="已人工确认", creator="domain-admin", confirmed_by="domain-admin",
                    confirmed_at=datetime.now(timezone.utc),
                    confirmed_content_hash=hashlib.sha256(content.encode("utf-8")).hexdigest(),
                )
                db.add(item)
                await db.flush()
                document_ids.append(item.id)
            await db.commit()
            record_id = record.id

        original_docx = storage._docx_bytes
        entered = threading.Event()
        release = threading.Event()

        def slow_docx(title, content):
            entered.set()
            release.wait(2)
            return original_docx(title, content)

        with patch.object(storage, "_docx_bytes", slow_docx):
            pending = asyncio.create_task(self.client.get(f"/api/v1/agent/documents/{document_ids[0]}/download"))
            self.assertTrue(await asyncio.to_thread(entered.wait, 2))
            await asyncio.sleep(0.01)
            self.assertFalse(pending.done(), "DOCX 下载生成不能阻塞事件循环")
            release.set()
            download = await pending
        self.assertEqual(download.status_code, 200, download.text)
        downloaded = Document(io.BytesIO(download.content))
        self.assertIn("正文第 0 份", [paragraph.text for paragraph in downloaded.paragraphs])

        entered.clear()
        release.clear()
        with patch.object(storage, "_docx_bytes", slow_docx):
            pending = asyncio.create_task(self.client.post(f"/api/v1/agent/documents/{document_ids[0]}/writeback"))
            self.assertTrue(await asyncio.to_thread(entered.wait, 2))
            await asyncio.sleep(0.01)
            self.assertFalse(pending.done(), "DOCX 回写生成不能阻塞事件循环")
            release.set()
            success = await pending
        self.assertEqual(success.status_code, 200, success.text)
        self.assertEqual(success.json()["record_id"], record_id)
        async with self.sessions() as db:
            attachment = await db.get(FileAttachment, success.json()["attachment_id"])
            self.assertEqual(attachment.category, "智能生成文书")
            self.assertEqual(attachment.size, Path(attachment.path).stat().st_size)
            self.assertIn("正文第 0 份", [paragraph.text for paragraph in Document(attachment.path).paragraphs])
            events = (await db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == record_id))).all()
            self.assertEqual([event.action for event in events], ["智能文档回写"])
        duplicate = await self.client.post(f"/api/v1/agent/documents/{document_ids[0]}/writeback")
        self.assertEqual(duplicate.status_code, 409, duplicate.text)

        original_write = Path.write_bytes
        partial_paths = []

        def fail_partial_write(path, content):
            if path.suffix == ".tmp":
                original_write(path, content[:16])
                partial_paths.append(path)
                raise OSError("injected partial DOCX write")
            return original_write(path, content)

        with patch.object(Path, "write_bytes", fail_partial_write):
            with self.assertRaisesRegex(OSError, "injected partial DOCX write"):
                await self.client.post(f"/api/v1/agent/documents/{document_ids[1]}/writeback")
        self.assertEqual(len(partial_paths), 1)
        self.assertFalse(partial_paths[0].exists())

        default_db = app.dependency_overrides[get_db]

        async def failing_commit_db():
            async with self.sessions() as db:
                with patch.object(db, "commit", side_effect=RuntimeError("injected writeback commit failure")):
                    yield db

        app.dependency_overrides[get_db] = failing_commit_db
        try:
            with self.assertRaisesRegex(RuntimeError, "injected writeback commit failure"):
                await self.client.post(f"/api/v1/agent/documents/{document_ids[2]}/writeback")
        finally:
            app.dependency_overrides[get_db] = default_db

        async def failing_commit_and_rollback_db():
            async with self.sessions() as db:
                with patch.object(db, "commit", side_effect=RuntimeError("injected writeback commit failure")):
                    with patch.object(db, "rollback", side_effect=RuntimeError("injected writeback rollback failure")):
                        yield db

        app.dependency_overrides[get_db] = failing_commit_and_rollback_db
        try:
            with self.assertRaises(BaseExceptionGroup) as failure:
                await self.client.post(f"/api/v1/agent/documents/{document_ids[3]}/writeback")
            self.assertEqual(
                [str(error) for error in failure.exception.exceptions],
                ["injected writeback commit failure", "injected writeback rollback failure"],
            )
        finally:
            app.dependency_overrides[get_db] = default_db

        async with self.sessions() as db:
            attachments = (await db.scalars(select(FileAttachment).where(FileAttachment.record_id == record_id))).all()
            events = (await db.scalars(select(WorkflowEvent).where(WorkflowEvent.record_id == record_id))).all()
            self.assertEqual(len(attachments), 1)
            self.assertEqual(len(events), 1)
        self.assertEqual(len(list(self.upload_root.glob("*.docx"))), 1)
        self.assertEqual(list(self.upload_root.glob("*.tmp")), [])


class SqliteDomainRouteStorageTest(DomainRouteStorageAssertions, unittest.IsolatedAsyncioTestCase):
    database_url = ""


@unittest.skipUnless(POSTGRES_URL, "需提供 AGENT_COMMANDS_TEST_POSTGRES_URL")
class PostgresDomainRouteStorageTest(DomainRouteStorageAssertions, unittest.IsolatedAsyncioTestCase):
    database_url = POSTGRES_URL
