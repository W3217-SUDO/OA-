"""问题反馈 BUG20261027002/003/005：付款申请单 Word、分配回款提交日期、银行导入页入库状态。"""

import io
import unittest
from datetime import datetime, timezone
from types import SimpleNamespace

from docx import Document

from app.areas.finance.incoming_payments import _fee_created_date
from app.areas.finance.incoming_query import matches_incoming_filters
from app.areas.finance.payment_print_word import payment_package_word


class FeeCreatedDateTest(unittest.TestCase):
    def test_uses_beijing_day_of_fee_creation(self):
        # UTC 10-26 17:30 是北京时间 10-27 01:30，日期按北京时间算。
        self.assertEqual(_fee_created_date(datetime(2026, 10, 26, 17, 30, tzinfo=timezone.utc)), "2026-10-27")
        self.assertEqual(_fee_created_date(datetime(2024, 10, 15, 9, 0)), "2024-10-15")
        self.assertEqual(_fee_created_date(None), "")


class BankImportStatusFilterTest(unittest.TestCase):
    def matches(self, status, value):
        item = SimpleNamespace(payer_name="上海恒达知识产权代理有限公司", status=status)
        return matches_incoming_filters(item, "finance-receipts-icbc", {"routeField3": value}, {}, {})

    def test_unclaimed_upload_is_not_imported_and_claimed_is_imported(self):
        self.assertTrue(self.matches("待认领", "未导入"))
        self.assertFalse(self.matches("待认领", "已导入"))
        for status in ("待分配", "部分分配", "已分配"):
            self.assertTrue(self.matches(status, "已导入"))
            self.assertFalse(self.matches(status, "未导入"))
        self.assertTrue(self.matches("已分配", "请选择"))


class PaymentPackageWordTest(unittest.TestCase):
    def test_word_follows_payment_request_template(self):
        document_data = {"data": {
            "total_amount": 100, "payee": "上海恒达知识产权代理有限公司", "account_bank": "工行徐汇淮海中路支行",
            "account": "1001071419100058177",
            "document_groups": [{
                "request_no": "P26102710453282068", "contract_no": "HZHT2510253", "contract_name": "测试11.13合同",
                "amount": 100, "remark": "",
                "items": [{"case_no": "HZMS2501006", "current_payment": 100, "fee_type": "一审诉讼费",
                           "fee_remark": "法院号:无", "applicant": "审核管理（赵媛）", "payer": "测试11.13"}],
            }],
        }}
        filename, content = payment_package_word(document_data, "P261027-75992859")
        self.assertEqual(filename, "P261027-75992859-付款申请单.docx")
        document = Document(io.BytesIO(content))
        section = document.sections[0]
        self.assertGreater(section.page_width, section.page_height)
        texts = {cell.text for row in document.tables[0].rows for cell in row.cells}
        for expected in ("上海沪知律师事务所　付款申请单", "P261027-75992859", "￥100.00", "上海恒达知识产权代理有限公司",
                         "工行徐汇淮海中路支行", "1001071419100058177", "P26102710453282068", "HZHT2510253",
                         "HZMS2501006", "一审诉讼费", "法院号:无", "审核管理（赵媛）", "测试11.13", "小计"):
            self.assertIn(expected, texts)
        self.assertIn("出纳签字", document.paragraphs[-1].text)


if __name__ == "__main__":
    unittest.main()
