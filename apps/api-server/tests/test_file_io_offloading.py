"""验证附件预览、正式发文 ZIP 与盖章文件的线程和持久化边界。"""

import asyncio
import io
import os
import threading
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import pypdfium2 as pdfium
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool
from starlette.concurrency import run_in_threadpool
from xlwt import Workbook as LegacyWorkbook

from app.agent_attachment_reader import read_attachment
from app.areas.aws.router import (
    download_official_outgoing_documents, upload_official_outgoing_stamp_file,
)
from app.areas.system.router import (
    get_pdf_preview_metadata, preview_attachment, render_pdf_preview_page,
)
from app.core.constants import UPLOAD_ROOT
from app.core.bank_document_recognition import _close_pages, _next_page, document_pages
from app.core.storage import _copy_seal_source_attachments, _render_pdf_preview_page
from app.database import Base
from app.models import BusinessRecord, FileAttachment, OfficialOutgoingDocument, WorkflowEvent
from app.models_shared import OfficialOutgoingBatchInput


POSTGRES_URL = os.environ.get("FILE_IO_TEST_POSTGRES_URL", "")


class FileIoAssertions:
    async def asyncSetUp(self):
        if self.database_url:
            if make_url(self.database_url).database != "oa_test_enterprise_commands":
                raise RuntimeError("文件 I/O 测试只能使用专用 PostgreSQL 测试库")
            self.schema = f"oa_test_enterprise_files_{uuid4().hex[:12]}"
            self.admin_engine = create_async_engine(self.database_url)
            async with self.admin_engine.begin() as connection:
                await connection.execute(text(f"CREATE SCHEMA {self.schema}"))
            self.engine = create_async_engine(
                self.database_url,
                connect_args={"server_settings": {"search_path": self.schema}},
            )
        else:
            self.engine = create_async_engine(
                "sqlite+aiosqlite:///:memory:",
                connect_args={"check_same_thread": False},
                poolclass=StaticPool,
            )
        async with self.engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
        self.created_paths = []
        self.identity = {"username": "file-io-admin", "role": "admin", "_actual_role_ids": ["admin"]}

    async def asyncTearDown(self):
        await self.engine.dispose()
        if self.database_url:
            async with self.admin_engine.begin() as connection:
                await connection.execute(text(f"DROP SCHEMA {self.schema} CASCADE"))
            await self.admin_engine.dispose()
        for path in self.created_paths:
            path.unlink(missing_ok=True)

    def _path(self, suffix):
        path = UPLOAD_ROOT / f"oa-test-enterprise-file-{uuid4().hex}{suffix}"
        self.created_paths.append(path)
        return path

    async def _attachment(self, db, path, original_name, content_type):
        item = FileAttachment(
            original_name=original_name, stored_name=path.name, content_type=content_type,
            size=path.stat().st_size, path=str(path), uploader=self.identity["username"],
        )
        db.add(item)
        await db.commit()
        return item.id

    async def test_workbook_and_pdf_preview(self):
        xlsx_path = self._path(".xlsx")
        workbook = Workbook()
        workbook.active["A1"] = "线程中读取的表格"
        workbook.save(xlsx_path)
        xls_path = self._path(".xls")
        legacy = LegacyWorkbook()
        legacy.add_sheet("第一页").write(0, 0, "旧表格")
        legacy.save(str(xls_path))
        pdf_path = self._path(".pdf")
        document = pdfium.PdfDocument.new()
        document.new_page(612, 792).close()
        document.save(pdf_path)
        document.close()
        broken_path = self._path(".pdf")
        broken_path.write_bytes(b"invalid PDF")

        async with self.sessions() as db:
            xlsx_id = await self._attachment(db, xlsx_path, "表格.xlsx", "application/octet-stream")
            xls_id = await self._attachment(db, xls_path, "旧表格.xls", "application/octet-stream")
            pdf_id = await self._attachment(db, pdf_path, "预览.pdf", "application/pdf")
            broken_id = await self._attachment(db, broken_path, "损坏.pdf", "application/pdf")

        async with self.sessions() as db:
            xlsx = await preview_attachment(xlsx_id, None, self.identity, db)
            xls = await preview_attachment(xls_id, None, self.identity, db)
            metadata = await get_pdf_preview_metadata(pdf_id, None, self.identity, db)
            image = await render_pdf_preview_page(pdf_id, 1, 900, None, self.identity, db)
            with self.assertRaises(HTTPException) as invalid_page:
                await render_pdf_preview_page(pdf_id, 2, 900, None, self.identity, db)
            with self.assertRaises(HTTPException) as invalid_file:
                await get_pdf_preview_metadata(broken_id, None, self.identity, db)
        self.assertIn("线程中读取的表格", xlsx["text"])
        self.assertIn("旧表格", str(xls["sheets"]))
        self.assertEqual(metadata.body.count(b'"page_count":1'), 1)
        self.assertEqual(image.body[:8], b"\x89PNG\r\n\x1a\n")
        self.assertEqual(invalid_page.exception.status_code, 404)
        self.assertEqual(invalid_file.exception.status_code, 422)

        from app.core import storage

        original_render = storage._render_pdf_preview_page
        entered = threading.Event()
        finished = threading.Event()

        def delayed_render(path, page_number, width):
            entered.set()
            time.sleep(0.2)
            try:
                return original_render(path, page_number, width)
            finally:
                finished.set()

        async with self.sessions() as db:
            with patch.object(storage, "_render_pdf_preview_page", delayed_render):
                request = asyncio.create_task(render_pdf_preview_page(pdf_id, 1, 900, None, self.identity, db))
                self.assertTrue(await asyncio.to_thread(entered.wait, 2))
                await asyncio.sleep(0.01)
                self.assertFalse(finished.is_set(), "PDF 渲染期间事件循环应继续执行")
                self.assertEqual((await request).body[:8], b"\x89PNG\r\n\x1a\n")

    async def test_zip_content_and_stamp_rollback(self):
        source = self._path(".txt")
        source.write_bytes(b"official outgoing source")
        async with self.sessions() as db:
            record = BusinessRecord(
                module="official_outgoing", serial_no="ENTERPRISE-OUTGOING-1",
                title="正式发文", customer="测试客户", status="已通过",
                owner=self.identity["username"], department="上海", data={},
            )
            db.add(record)
            await db.flush()
            detail = OfficialOutgoingDocument(
                record_id=record.id, official_no=record.serial_no,
                source_type="case", created_by=self.identity["username"],
            )
            attachment = FileAttachment(
                record_id=record.id, original_name="来源.txt", stored_name=source.name,
                content_type="text/plain", size=source.stat().st_size, path=str(source),
                uploader=self.identity["username"],
            )
            db.add_all([detail, attachment])
            await db.commit()
            record_id = record.id
            attachment_id = attachment.id

        async with self.sessions() as db:
            response = await download_official_outgoing_documents(
                OfficialOutgoingBatchInput(record_ids=[record_id]), self.identity, db,
            )
            archive_bytes = b"".join([part async for part in response.body_iterator])
        with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
            self.assertEqual(archive.namelist(), [f"ENTERPRISE-OUTGOING-1/{attachment_id}-来源.txt"])
            self.assertEqual(archive.read(archive.namelist()[0]), b"official outgoing source")

        original_read = Path.read_bytes
        entered = threading.Event()
        finished = threading.Event()

        def delayed_read(path):
            if path == source:
                entered.set()
                time.sleep(0.2)
                finished.set()
            return original_read(path)

        async with self.sessions() as db:
            with patch.object(Path, "read_bytes", delayed_read):
                request = asyncio.create_task(download_official_outgoing_documents(
                    OfficialOutgoingBatchInput(record_ids=[record_id]), self.identity, db,
                ))
                self.assertTrue(await asyncio.to_thread(entered.wait, 2))
                await asyncio.sleep(0.01)
                self.assertFalse(finished.is_set(), "ZIP 读取期间事件循环应继续执行")
                await request

        source.unlink()
        async with self.sessions() as db:
            with self.assertRaises(HTTPException) as no_file:
                await download_official_outgoing_documents(
                    OfficialOutgoingBatchInput(record_ids=[record_id]), self.identity, db,
                )
        self.assertEqual(no_file.exception.status_code, 404)
        source.write_bytes(b"official outgoing source")

        baseline_files = set(UPLOAD_ROOT.iterdir())
        async with self.sessions() as db:
            with patch.object(db, "commit", side_effect=RuntimeError("模拟数据库提交失败")):
                with self.assertRaisesRegex(RuntimeError, "模拟数据库提交失败"):
                    await upload_official_outgoing_stamp_file(
                        record_id, UploadFile(filename="盖章.pdf", file=io.BytesIO(b"stamp")),
                        self.identity, db,
                    )
        self.assertEqual(set(UPLOAD_ROOT.iterdir()), baseline_files)
        async with self.sessions() as db:
            stored_record = await db.get(BusinessRecord, record_id)
            stored_detail = await db.scalar(select(OfficialOutgoingDocument).where(
                OfficialOutgoingDocument.record_id == record_id,
            ))
            file_count = await db.scalar(select(func.count(FileAttachment.id)).where(
                FileAttachment.record_id == record_id,
            ))
            event_count = await db.scalar(select(func.count(WorkflowEvent.id)).where(
                WorkflowEvent.record_id == record_id,
            ))
        self.assertEqual(stored_record.status, "已通过")
        self.assertIsNone(stored_detail.stamp_attachment_id)
        self.assertEqual(file_count, 1)
        self.assertEqual(event_count, 0)

        async with self.sessions() as db:
            response = await upload_official_outgoing_stamp_file(
                record_id, UploadFile(filename="盖章.pdf", file=io.BytesIO(b"stamp")),
                self.identity, db,
            )
        self.assertEqual(response["status"], "已盖章")
        async with self.sessions() as db:
            stored_detail = await db.scalar(select(OfficialOutgoingDocument).where(
                OfficialOutgoingDocument.record_id == record_id,
            ))
            stamped_file = await db.get(FileAttachment, stored_detail.stamp_attachment_id)
            event_count = await db.scalar(select(func.count(WorkflowEvent.id)).where(
                WorkflowEvent.record_id == record_id,
            ))
        self.created_paths.append(Path(stamped_file.path))
        self.assertEqual(Path(stamped_file.path).read_bytes(), b"stamp")
        self.assertEqual(event_count, 1)

    async def test_seal_source_copy_and_partial_file_cleanup(self):
        source_path = self._path(".pdf")
        source_path.write_bytes(b"source content")
        async with self.sessions() as db:
            customer = BusinessRecord(
                module="customer", serial_no="ENTERPRISE-COPY-CUSTOMER",
                title="来源客户", customer="", status="正常",
                owner=self.identity["username"], department="上海", data={},
            )
            first_seal = BusinessRecord(
                module="seal_application", serial_no="ENTERPRISE-SEAL-1",
                title="第一份用印", customer="来源客户", status="草稿",
                owner=self.identity["username"], department="上海", data={},
            )
            second_seal = BusinessRecord(
                module="seal_application", serial_no="ENTERPRISE-SEAL-2",
                title="第二份用印", customer="来源客户", status="草稿",
                owner=self.identity["username"], department="上海", data={},
            )
            db.add_all([customer, first_seal, second_seal])
            await db.flush()
            original = FileAttachment(
                record_id=customer.id, original_name="来源.pdf", stored_name=source_path.name,
                content_type="application/pdf", size=source_path.stat().st_size,
                path=str(source_path), uploader=self.identity["username"],
            )
            db.add(original)
            await db.commit()
            first_id, second_id, original_id = first_seal.id, second_seal.id, original.id

        async with self.sessions() as db:
            seal = await db.get(BusinessRecord, first_id)
            copied_paths = await _copy_seal_source_attachments(
                seal, [original_id], self.identity, db,
            )
            await db.commit()
        self.created_paths.extend(copied_paths)
        self.assertEqual(len(copied_paths), 1)
        self.assertEqual(copied_paths[0].read_bytes(), b"source content")
        async with self.sessions() as db:
            seal = await db.get(BusinessRecord, first_id)
            self.assertEqual(seal.data["document_names"], "来源.pdf")
            self.assertEqual(await db.scalar(select(func.count(FileAttachment.id)).where(
                FileAttachment.record_id == first_id,
            )), 1)

        baseline_files = set(UPLOAD_ROOT.iterdir())

        def incomplete_copy(_source, target):
            Path(target).write_bytes(b"partial")
            raise OSError("模拟复制中断")

        from app.core import storage

        async with self.sessions() as db:
            seal = await db.get(BusinessRecord, second_id)
            with patch.object(storage, "copyfile", incomplete_copy):
                with self.assertRaisesRegex(OSError, "模拟复制中断"):
                    await _copy_seal_source_attachments(seal, [original_id], self.identity, db)
        self.assertEqual(set(UPLOAD_ROOT.iterdir()), baseline_files)
        async with self.sessions() as db:
            self.assertEqual(await db.scalar(select(func.count(FileAttachment.id)).where(
                FileAttachment.record_id == second_id,
            )), 0)

    async def test_pdfium_is_serialized_across_three_readers(self):
        path = self._path(".pdf")
        document = pdfium.PdfDocument.new()
        document.new_page(612, 792).close()
        document.save(path)
        document.close()
        raw = path.read_bytes()

        original_render = pdfium.PdfPage.render
        counter_lock = threading.Lock()
        active = 0
        maximum = 0
        calls = 0

        def measured_render(page, *args, **kwargs):
            nonlocal active, maximum, calls
            with counter_lock:
                active += 1
                calls += 1
                maximum = max(maximum, active)
            try:
                time.sleep(0.1)
                return original_render(page, *args, **kwargs)
            finally:
                with counter_lock:
                    active -= 1

        async def recognize_page():
            pages = document_pages(raw, "流水.pdf")
            try:
                return await run_in_threadpool(_next_page, pages)
            finally:
                await run_in_threadpool(_close_pages, pages)

        with patch.object(pdfium.PdfPage, "render", measured_render):
            preview, agent_reading, bank_page, invalid_page = await asyncio.gather(
                run_in_threadpool(_render_pdf_preview_page, path, 1, 900),
                asyncio.to_thread(read_attachment, path, path.name),
                recognize_page(),
                run_in_threadpool(_render_pdf_preview_page, path, 2, 900),
                return_exceptions=True,
            )
        self.assertEqual(preview[:8], b"\x89PNG\r\n\x1a\n")
        self.assertEqual(agent_reading.status, "visual")
        self.assertEqual(bank_page[0], "第1页")
        self.assertIsInstance(invalid_page, HTTPException)
        self.assertEqual(invalid_page.status_code, 404)
        self.assertEqual(calls, 3)
        self.assertEqual(maximum, 1)


class SQLiteFileIoTest(FileIoAssertions, unittest.IsolatedAsyncioTestCase):
    database_url = ""


@unittest.skipUnless(POSTGRES_URL, "设置 FILE_IO_TEST_POSTGRES_URL 运行独立 PostgreSQL 测试")
class PostgresFileIoTest(FileIoAssertions, unittest.IsolatedAsyncioTestCase):
    database_url = POSTGRES_URL
