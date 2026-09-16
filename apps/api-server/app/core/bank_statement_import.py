"""Parse each bank's native export layout, preserving source row numbers."""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile

ALIASES = {
    'payer_name': ('对方户名', '对方账户名称', '对方名称', '付款人', '付款方', '付款人名称', '付款户名', '付款账户名称', '回款单位', 'payer_name'),
    'bank_reference': ('银行流水号', '交易流水号', '业务编号', '流水号', '柜员交易号', '交易号', '交易序号', '银行交易流水号', '交易参考号', 'bank_reference'),
    'received_date': ('到账日期', '交易日期', '交易时间', '记账日期', '记账时间', '收付款日期', '回款日期', '入账日期', '入账时间', 'received_date'),
    'amount': ('到账金额', '贷方发生额', '贷方金额', '收入金额', '收入', '收款金额', '交易金额', '发生额', '金额', '回款金额', 'amount'),
    'direction': ('借贷标志', '借贷方向', '收付标志', '收付款标志', '交易方向'),
    'remark': ('摘要', '备注', '用途', '附言', 'remark'),
}
REQUIRED = ('payer_name', 'bank_reference', 'received_date', 'amount')

# Legacy AR/{bank}/PaymentService.Process: 1-based first row, 0-based columns.
# Explicit bank selection is authoritative; header synonyms must not override it.
BANK_TEMPLATES = {
    'icbc': {'name': '工行', 'first_row': 6, 'columns': {'payer_name': 5, 'amount': 6, 'received_date': 10, 'bank_reference': 13, 'direction': 8, 'remark': 12}},
    'citic': {'name': '中信', 'first_row': 15, 'columns': {'payer_name': 3, 'amount': 6, 'received_date': 0, 'bank_reference': 11, 'remark': 12}},
    'boc': {'name': '中行', 'first_row': 9, 'columns': {'payer_name': 5, 'amount': 13, 'received_date': 10, 'bank_reference': 17, 'remark': 25}},
}


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().lstrip("'\t")


def header_key(value):
    value = re.sub(r'[（(](?:元|人民币|人民币元|RMB|CNY)[）)]', '', cell_text(value), flags=re.I)
    return re.sub(r'[\s\u3000：:()（）\[\]]', '', value).casefold()


def statement_sheets(raw, filename):
    suffix = Path(filename).suffix.lower()
    if suffix in {'.csv', '.tsv'}:
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                content = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('CSV编码无法识别，请保存为UTF-8或GB18030')
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        yield 'CSV', csv.reader(io.StringIO(content), dialect)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 512 * 1024 * 1024:
                raise ValueError('Excel解压后超过512MB，请拆分文件')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                if sheet.max_column and sheet.max_column > 256:
                    raise ValueError('Excel列数超过256，请删除多余空白列')
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            book.close()
    elif suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=raw, on_demand=True)
        try:
            for sheet in book.sheets():
                def rows(current=sheet):
                    for index in range(current.nrows):
                        yield [xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                               if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                               for cell in current.row(index)]
                yield sheet.name, rows()
        finally:
            book.release_resources()
    else:
        raise ValueError('仅支持.xlsx、.xls或.csv银行流水文件')


def read_legacy_statement(raw, filename, bank_source):
    if bank_source not in BANK_TEMPLATES or Path(filename).suffix.lower() not in {'.xlsx', '.xls'}:
        return None
    template = BANK_TEMPLATES[bank_source]
    first_row, mapping = template['first_row'], template['columns']
    result = []
    sheets = statement_sheets(raw, filename)
    try:
        sheet_name, rows = next(sheets)
        for row_number, cells in enumerate(rows, 1):
            if row_number > 50000:
                raise ValueError('单个工作表超过50000行，请拆分导入')
            if row_number < first_row or len(cells) < 2 or not cell_text(cells[1]):
                continue
            if any(cell_text(cell) in {'合计', '总计', '小计'} for cell in cells[:2]):
                continue
            row = {field: cells[index] if index < len(cells) else '' for field, index in mapping.items()}
            # Field validation belongs to the per-row API result, not the entire file.
            result.append((f"{template['name']} / {sheet_name}", row_number, row))
    finally:
        sheets.close()
    if not result:
        raise ValueError(f"{template['name']}导出模板从第{first_row}行读取，未找到流水明细；请在对应银行入口上传完整导出表")
    return result


def read_statement(raw, filename, bank_source=''):
    native = read_legacy_statement(raw, filename, bank_source)
    if native is not None:
        return native
    result = []
    matched_sheets = 0
    for sheet_name, rows in statement_sheets(raw, filename):
        mapping = None
        for row_number, cells in enumerate(rows, 1):
            if row_number > 50000:
                raise ValueError('单个工作表超过50000行，请拆分导入')
            if not any(cell_text(cell) for cell in cells):
                continue
            headers = [header_key(cell) for cell in cells]
            detected = {}
            for field, aliases in ALIASES.items():
                for alias in aliases:
                    if header_key(alias) in headers:
                        detected[field] = headers.index(header_key(alias))
                        break
            if all(field in detected for field in REQUIRED):
                mapping = detected
                matched_sheets += 1
                continue
            if mapping is None:
                if row_number >= 100:
                    break
                continue
            row = {field: cells[index] if index < len(cells) else '' for field, index in mapping.items()}
            # Totals are not transaction records; other malformed rows are reported by the caller.
            if any(cell_text(cell) in {'合计', '总计', '小计'} for cell in cells[:2]):
                continue
            result.append((sheet_name, row_number, row))
    if not matched_sheets:
        raise ValueError('未识别银行流水表头，需要对方户名、银行流水号、到账日期及到账金额列；请上传银行明细导出文件')
    if not result:
        raise ValueError('文件没有可导入的流水明细')
    return result


def parse_received_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value).replace('/', '-').replace('.', '-').replace('年', '-').replace('月', '-').replace('日', '')
    for pattern in ('%Y-%m-%d', '%Y%m%d'):
        try:
            return datetime.strptime(re.split('[ T]', text)[0], pattern).date()
        except ValueError:
            pass
    raise ValueError('到账日期无法识别，请使用年月日日期')


def parse_amount(value):
    try:
        amount = Decimal(cell_text(value).replace(',', '').replace('，', '').replace('￥', '').replace('¥', ''))
    except InvalidOperation as exc:
        raise ValueError('到账金额不是有效数字') from exc
    if not amount.is_finite() or amount <= 0 or amount >= Decimal('1000000000000000'):
        raise ValueError('到账金额必须为大于0的有效数字')
    return float(amount.quantize(Decimal('0.01')))
