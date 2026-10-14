"""个人文档的受限 Markdown 与 DOCX 格式往返，不参与通用文本提取。"""

from io import BytesIO
from string import punctuation

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from markdown_it import MarkdownIt


_MARKDOWN = MarkdownIt("commonmark", {"html": False})
_FORMAT_ERROR = "当前在线编辑仅支持一级至三级标题、粗体、斜体及一级普通列表；原文件未修改"


def _escaped_text(value: str) -> str:
    return "".join("\\" + char if char in punctuation else char for char in value)


def _paragraph_markdown(paragraph) -> str:
    if any(child.tag.rsplit("}", 1)[-1] not in {"pPr", "r", "bookmarkStart", "bookmarkEnd"} for child in paragraph._p):
        raise ValueError(_FORMAT_ERROR)
    groups: list[tuple[str, bool, bool]] = []
    for run in paragraph.runs:
        if any(child.tag.rsplit("}", 1)[-1] not in {"rPr", "t", "tab", "br", "cr"} for child in run._r):
            raise ValueError(_FORMAT_ERROR)
        bold, italic = run.bold, run.italic
        style = run.style
        while style is not None and (bold is None or italic is None):
            if bold is None:
                bold = style.font.bold
            if italic is None:
                italic = style.font.italic
            style = style.base_style
        flags = bool(bold), bool(italic)
        if groups and groups[-1][1:] == flags:
            previous = groups.pop()
            groups.append((previous[0] + run.text, *flags))
        else:
            groups.append((run.text, *flags))
    fragments = []
    for text, bold, italic in groups:
        marker = "***" if bold and italic else "**" if bold else "*" if italic else ""
        lines = []
        for line in text.split("\n"):
            core = line.strip()
            if not core:
                lines.append(line)
                continue
            leading = line[:len(line) - len(line.lstrip())]
            trailing = line[len(line.rstrip()):]
            lines.append(leading + marker + _escaped_text(core) + marker + trailing)
        fragments.append("  \n".join(lines))
    return "".join(fragments)


def _paragraph_numbering(document, paragraph):
    properties = paragraph._p.pPr
    numbering = properties.numPr if properties is not None else None
    style = paragraph.style
    while numbering is None and style is not None:
        properties = style.element.pPr
        numbering = properties.numPr if properties is not None else None
        style = style.base_style
    if numbering is None or numbering.numId is None:
        return None
    if numbering.ilvl is not None and numbering.ilvl.val != 0:
        raise ValueError(_FORMAT_ERROR)
    num_id = numbering.numId.val
    root = document.part.numbering_part.element
    definitions = root.xpath(f'w:num[@w:numId="{num_id}"]')
    if len(definitions) != 1:
        raise ValueError("文档列表编号定义缺失，原文件未修改")
    reference = definitions[0].xpath("w:abstractNumId")
    if len(reference) != 1:
        raise ValueError("文档列表编号定义缺失，原文件未修改")
    abstract_id = int(reference[0].get(qn("w:val")))
    levels = root.xpath(f'w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="0"]')
    if len(levels) != 1:
        raise ValueError(_FORMAT_ERROR)
    formats = levels[0].findall(qn("w:numFmt"))
    if len(formats) != 1 or formats[0].get(qn("w:val")) not in {"decimal", "bullet"}:
        raise ValueError(_FORMAT_ERROR)
    starts = definitions[0].xpath('w:lvlOverride[@w:ilvl="0"]/w:startOverride') or levels[0].findall(qn("w:start"))
    start = int(starts[0].get(qn("w:val"))) if starts else 1
    return num_id, formats[0].get(qn("w:val")) == "decimal", start


def personal_docx_markdown(document, *, generated_title: str) -> str:
    if document.tables:
        raise ValueError("含表格的文档不支持本次受限格式编辑，原文件未修改")
    lines: list[str] = []
    counters: dict[int, int] = {}
    for index, paragraph in enumerate(document.paragraphs):
        style_id = paragraph.style.style_id
        if style_id.startswith("Heading") and style_id not in {"Heading1", "Heading2", "Heading3"}:
            raise ValueError(_FORMAT_ERROR)
        if style_id in {"ListBullet2", "ListBullet3", "ListNumber2", "ListNumber3"}:
            raise ValueError(_FORMAT_ERROR)
        # 写入器额外生成的文件标题不是正文，重存时不可重复追加。
        if index == 0 and style_id == "Title" and paragraph.text == generated_title:
            continue
        value = _paragraph_markdown(paragraph)
        if style_id in {"Heading1", "Heading2", "Heading3"}:
            value = "#" * int(style_id[-1]) + " " + value
        else:
            numbering = _paragraph_numbering(document, paragraph)
            if numbering is not None:
                num_id, ordered, start = numbering
                counters[num_id] = counters[num_id] + 1 if num_id in counters else start
                value = (f"{counters[num_id]}. " if ordered else "- ") + value
        lines.append(value)
    return "\n\n".join(lines).strip()


def _new_numbering(document, *, ordered: bool, start: int) -> int:
    root = document.part.numbering_part.element
    abstract_id = max((int(item.get(qn("w:abstractNumId"))) for item in root.xpath("w:abstractNum")), default=-1) + 1
    num_id = max((int(item.get(qn("w:numId"))) for item in root.xpath("w:num")), default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", str(start)), ("numFmt", "decimal" if ordered else "bullet"), ("lvlText", "%1." if ordered else "\u2022")):
        element = OxmlElement("w:" + tag); element.set(qn("w:val"), value); level.append(element)
    abstract.append(level); root.append(abstract)
    number = OxmlElement("w:num"); number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId"); reference.set(qn("w:val"), str(abstract_id))
    number.append(reference); root.append(number)
    return num_id


def _add_inline(paragraph, children) -> None:
    bold = italic = 0
    for token in children:
        if token.type == "strong_open":
            bold += 1
        elif token.type == "strong_close":
            bold -= 1
        elif token.type == "em_open":
            italic += 1
        elif token.type == "em_close":
            italic -= 1
        elif token.type in {"text", "softbreak", "hardbreak"}:
            run = paragraph.add_run(token.content if token.type == "text" else " " if token.type == "softbreak" else "\n")
            if bold:
                run.bold = True
            if italic:
                run.italic = True
        else:
            raise ValueError(_FORMAT_ERROR)


def personal_markdown_docx(title: str, content: str) -> bytes:
    tokens = _MARKDOWN.parse(content)
    document = Document()
    document.add_heading(title, level=0)
    paragraph = None
    current_list = None
    pending_item = False
    for token in tokens:
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            if current_list is not None:
                raise ValueError(_FORMAT_ERROR)
            ordered = token.type == "ordered_list_open"
            start = token.attrGet("start")
            current_list = (ordered, _new_numbering(document, ordered=ordered, start=int(start) if start is not None else 1))
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            current_list = None
        elif token.type == "list_item_open":
            pending_item = True
        elif token.type == "list_item_close":
            pending_item = False
        elif token.type == "heading_open":
            if token.tag not in {"h1", "h2", "h3"} or current_list is not None:
                raise ValueError(_FORMAT_ERROR)
            paragraph = document.add_heading(level=int(token.tag[1]))
        elif token.type == "paragraph_open":
            if current_list is None:
                paragraph = document.add_paragraph()
            else:
                if not pending_item:
                    raise ValueError(_FORMAT_ERROR)
                ordered, num_id = current_list
                paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
                numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                numbering.get_or_add_ilvl().val = 0
                numbering.get_or_add_numId().val = num_id
                pending_item = False
        elif token.type == "inline":
            if paragraph is None:
                raise ValueError(_FORMAT_ERROR)
            _add_inline(paragraph, token.children or [])
        elif token.type in {"paragraph_close", "heading_close"}:
            paragraph = None
        else:
            raise ValueError(_FORMAT_ERROR)
    output = BytesIO(); document.save(output)
    return output.getvalue()
