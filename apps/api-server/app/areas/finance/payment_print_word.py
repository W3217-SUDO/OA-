"""普通付款包的“付款申请单”Word：版式按原系统“提交并打印”导出的模板（横向 A4、六列表格、底部三处签字）。"""
from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

_GRAY = "D9D9D9"


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _money(value) -> str:
    try:
        return f"￥{float(value or 0):.2f}"
    except (TypeError, ValueError):
        return "￥0.00"


def _fill(cell, text: str, *, bold: bool = False, shade: bool = False, size: float = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if shade:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), _GRAY)
        properties.append(shading)


def _row(table, values: list[str], *, spans: list[int] | None = None, bold: bool = False, shade: bool = False) -> None:
    """按跨列数写一行；spans 合计必须等于 6。"""
    spans = spans or [1] * len(values)
    cells = table.add_row().cells
    start = 0
    for value, span in zip(values, spans):
        cell = cells[start]
        if span > 1:
            cell = cell.merge(cells[start + span - 1])
        _fill(cell, value, bold=bold, shade=shade)
        start += span


def payment_package_word(document_data: dict, package_no: str) -> tuple[str, bytes]:
    data = document_data.get("data") or {}
    groups = data.get("document_groups") or []
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.2))

    table = document.add_table(rows=0, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_cells = table.add_row().cells
    title = title_cells[0].merge(title_cells[5])
    _fill(title, f"{_text(data.get('company_name')) or '上海沪知律师事务所'}　付款申请单", bold=True, size=14)
    _row(table, ["打包流水号：", package_no, "打印日期：", f"{datetime.now():%Y-%m-%d}", "付款总金额：",
                 _money(data.get("total_amount") or data.get("amount"))])
    _row(table, ["收款单位：", _text(data.get("payee_display_name") or data.get("payee"))], spans=[1, 5])
    _row(table, ["开户行：", _text(data.get("account_bank") or data.get("bank_name") or data.get("payee_bank") or data.get("bank"))], spans=[1, 5])
    _row(table, ["银行账号：", _text(data.get("account") or data.get("bank_account") or data.get("payee_account"))], spans=[1, 5])
    for group in groups:
        _row(table, ["请票单号：", _text(group.get("request_no")), "合同编号：", _text(group.get("contract_no")),
                     "合同名称：", _text(group.get("contract_name"))], shade=True)
        _row(table, ["案号", "付款金额", "费用类型", "备注", "申请人", "交款人"], shade=True)
        for item in group.get("items") or []:
            _row(table, [
                _text(item.get("case_no")),
                _money(item.get("current_payment") if item.get("current_payment") is not None else item.get("amount")),
                _text(item.get("expense_subtype") or item.get("fee_type")),
                _text(item.get("fee_remark")),
                _text(item.get("applicant_display_name") or item.get("applicant") or group.get("applicant")),
                _text(item.get("payer_display_name") or item.get("payer") or group.get("payer")),
            ])
        _row(table, ["备注：", _text(group.get("remark") or group.get("description"))], spans=[1, 5])
        _row(table, ["小计", _money(group.get("amount"))], spans=[2, 4])
    signature = document.add_paragraph()
    signature.add_run("案件管理人签字：" + "　" * 14 + "审批人签字：" + "　" * 14 + "出纳签字：").font.size = Pt(10)
    buffer = BytesIO()
    document.save(buffer)
    return f"{package_no}-付款申请单.docx", buffer.getvalue()
