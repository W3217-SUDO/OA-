"""只读接口等价性快照工具。

用途：在改动业务代码前后，用同一份用户身份与参数直接调用路由处理函数，
把**全部分页拼接**后的返回值规范化写出，供 compare_snapshots.py 逐字段比较。

约束：
- 数据库连接强制只读事务；本工具不写任何业务数据。
- 不经过 HTTP、不经过 FastAPI 依赖注入、不经过共享查询缓存。
- 参数必须显式传入（路由函数的默认值是 FieldInfo 对象，不能依赖默认值）。

运行示例（生产或本地）：
    PYTHONPATH=/path/to/api python -B scripts/perf/snapshot_endpoints.py \
        --cases scripts/perf/snapshot_cases.json --out /tmp/oa-perf/before \
        --user admin --only finance.internal_fees
"""
from __future__ import annotations

import argparse
import asyncio
import importlib
import inspect
import json
import sys
import time
import traceback
from collections import defaultdict
from contextlib import asynccontextmanager
from pathlib import Path


from fastapi import Request
from sqlalchemy import event, select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.config import settings
from app.models import User


# ---------------------------------------------------------------- 规范化

def _normalize(value):
    """把返回值转成可稳定比较的结构；集合排序，日期转字符串。"""
    if isinstance(value, dict):
        return {str(key): _normalize(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_normalize(item) for item in value]
    if isinstance(value, (set, frozenset)):
        return sorted((_normalize(item) for item in value), key=lambda item: (type(item).__name__, str(item)))
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _dumps(value) -> str:
    return json.dumps(_normalize(value), ensure_ascii=False, sort_keys=True,
                      separators=(",", ":"), allow_nan=False)


# ---------------------------------------------------------------- 计时

class QueryCounter:
    def __init__(self):
        self.statements: list[tuple[float, str, int]] = []

    def install(self, engine):
        @event.listens_for(engine.sync_engine, "before_cursor_execute")
        def _before(conn, cursor, statement, parameters, context, executemany):
            context._perf_started = time.perf_counter()

        @event.listens_for(engine.sync_engine, "after_cursor_execute")
        def _after(conn, cursor, statement, parameters, context, executemany):
            started = getattr(context, "_perf_started", None)
            if started is None:
                return
            count = len(parameters) if isinstance(parameters, (list, tuple, dict)) else 0
            self.statements.append((time.perf_counter() - started, " ".join(statement.split()), count))


# ---------------------------------------------------------------- 身份

async def build_identity(db: AsyncSession, username: str, path: str) -> dict:
    """按 app/security.py 的 current_identity 口径构造身份字典。"""
    from app.core.permissions import _user_permission_payload
    from app.security import _page_menu_authorized, user_role_ids

    user = await db.scalar(select(User).where(User.username == username))
    if not user or not user.is_active:
        raise SystemExit(f"用户不可用：{username}")
    role_ids = user_role_ids(user)
    permission = await _user_permission_payload(user, db)
    # 页面能力只影响管理员降级，快照按真实路径计算，保证与线上实际口径一致。
    scope = {
        "type": "http",
        "method": "GET",
        "path": path,
        "raw_path": path.encode(),
        "query_string": b"",
        "headers": [],
        "scheme": "http",
        "server": ("127.0.0.1", 8001),
        "client": ("127.0.0.1", 0),
        "root_path": "",
    }
    request = Request(scope)
    page_capability = await _page_menu_authorized(request, db, permission.get("menu_keys") or [])
    effective_role_ids = ["admin", *role_ids] if page_capability and "admin" not in role_ids else role_ids
    effective_action_keys = ["*"] if page_capability else permission.get("action_keys") or []
    return {
        "username": user.username,
        "role": "admin" if page_capability else role_ids[0],
        "role_ids": effective_role_ids,
        "display_name": user.display_name,
        "department": user.department,
        "permission_role": (user.profile or {}).get("permission_role") or "",
        "staff_role": (user.profile or {}).get("staff_role") or "",
        "position": (user.profile or {}).get("position") or "",
        "must_change_password": user.must_change_password,
        "menu_keys": permission.get("menu_keys") or [],
        "action_keys": effective_action_keys,
        "data_scope": permission.get("data_scope"),
        "_actual_role": role_ids[0],
        "_actual_role_ids": role_ids,
        "_page_menu_capability": page_capability,
        "_request_path": path,
    }


# ---------------------------------------------------------------- 返回值读取

async def read_response(result):
    """导出接口可能返回 Response/StreamingResponse；统一转成可比较内容。"""
    from starlette.responses import Response, StreamingResponse

    if isinstance(result, StreamingResponse):
        chunks = []
        async for chunk in result.body_iterator:
            chunks.append(chunk if isinstance(chunk, bytes) else str(chunk).encode())
        payload = b"".join(chunks)
        return _file_content(result.media_type or "", payload)
    if isinstance(result, Response):
        return _file_content(result.media_type or "", result.body)
    return result


def _file_content(media_type: str, payload: bytes):
    if not payload:
        return {"kind": "empty"}
    if "spreadsheet" in media_type or payload[:2] == b"PK":
        try:
            from io import BytesIO

            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(payload), data_only=True)
            return {"kind": "xlsx", "sheets": {
                sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
                for sheet in workbook.worksheets
            }}
        except Exception as error:  # noqa: BLE001 - 退回摘要，绝不吞掉差异
            return {"kind": "xlsx-unparsed", "error": type(error).__name__, "size": len(payload)}
    if "csv" in media_type or media_type.startswith("text/"):
        return {"kind": "text", "text": payload.decode("utf-8", "replace")}
    return {"kind": "binary", "size": len(payload)}


# ---------------------------------------------------------------- 用例执行

async def call_endpoint(module_name: str, function_name: str, params: dict, user: str, path: str,
                        db_factory):
    """调用路由处理函数；分页接口循环抓取全部分页。"""
    module = importlib.import_module(module_name)
    function = getattr(module, function_name)
    signature = inspect.signature(function)

    async with db_factory() as db:
        identity = await build_identity(db, user, path)

    page = int(params.get("page") or 1)
    page_size = int(params.get("page_size") or 200)
    collected = []
    totals = {}
    total = None
    pages = 0
    while True:
        arguments = dict(params)
        arguments["identity"] = identity
        if "page" in signature.parameters:
            arguments["page"] = page
        if "page_size" in signature.parameters:
            arguments["page_size"] = page_size
        unsupported = [name for name in arguments if name not in signature.parameters]
        if unsupported:
            raise SystemExit(f"{module_name}.{function_name} 不接受参数：{unsupported}")
        async with db_factory() as db:
            result = await function(**arguments, db=db)
        result = await read_response(result)
        await db.close() if False else None
        if isinstance(result, dict) and "items" in result:
            collected.extend(result.get("items") or [])
            for key, value in result.items():
                if key in {"items", "page", "page_size"}:
                    continue
                totals[key] = value
            total = result.get("total")
            pages += 1
            if not result.get("items") or len(collected) >= int(total or 0) or pages > 500:
                break
            page += 1
            continue
        return result
    return {**totals, "items": collected, "pages_fetched": pages}


# ---------------------------------------------------------------- 主流程

def make_factory(engine):
    @asynccontextmanager
    async def factory():
        async with AsyncSession(engine, expire_on_commit=False) as session:
            try:
                yield session
            finally:
                await session.rollback()
    return factory


async def run(args) -> int:
    cases = json.loads(Path(args.cases).read_text(encoding="utf-8"))
    selected = [case for case in cases if not args.only or case["id"] in set(args.only)]
    if not selected:
        raise SystemExit("没有匹配的用例")

    engine = create_async_engine(
        settings.database_url,
        connect_args={"server_settings": {
            "default_transaction_read_only": "on",
            "statement_timeout": args.timeout_ms,
        }},
        pool_size=1,
        max_overflow=0,
    )
    counter = QueryCounter()
    counter.install(engine)
    factory = make_factory(engine)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    timing = {}
    failures = 0

    for case in selected:
        params = {key: value for key, value in (case.get("params") or {}).items()}
        if args.page_size:
            params["page_size"] = args.page_size
        user = args.user or case.get("user") or "admin"
        counter.statements.clear()
        started = time.perf_counter()
        try:
            payload = await call_endpoint(case["endpoint"].rsplit(".", 1)[0], case["endpoint"].rsplit(".", 1)[1],
                                          params, user, case.get("path") or "/", factory)
            elapsed = time.perf_counter() - started
            (out / f"{case['id']}.json").write_text(_dumps(payload), encoding="utf-8")
            timing[case["id"]] = {
                "user": user,
                "wall_ms": round(elapsed * 1000, 1),
                "sql_ms": round(sum(item[0] for item in counter.statements) * 1000, 1),
                "statements": len(counter.statements),
                "ok": True,
            }
            print(f"OK   {case['id']:<46} {elapsed*1000:8.0f}ms sql={timing[case['id']]['sql_ms']:8.0f}ms "
                  f"stmts={len(counter.statements)}", flush=True)
        except Exception as error:  # noqa: BLE001 - 失败必须记录而不是中断整批
            failures += 1
            elapsed = time.perf_counter() - started
            timing[case["id"]] = {
                "user": user,
                "wall_ms": round(elapsed * 1000, 1),
                "statements": len(counter.statements),
                "ok": False,
                "error": f"{type(error).__name__}: {error}",
                "trace": traceback.format_exc().splitlines()[-6:],
            }
            (out / f"{case['id']}.json").write_text(_dumps({"error": timing[case["id"]]["error"]}), encoding="utf-8")
            print(f"FAIL {case['id']:<46} {type(error).__name__}: {error}", flush=True)
        if args.verbose and counter.statements:
            groups = defaultdict(lambda: [0, 0.0, 0])
            for duration, statement, count in counter.statements:
                entry = groups[statement[:150]]
                entry[0] += 1
                entry[1] += duration
                entry[2] = max(entry[2], count)
            for statement, (count, duration, params_count) in sorted(groups.items(), key=lambda item: -item[1][1])[:8]:
                print(f"       {duration*1000:8.0f}ms x{count:<4} params={params_count:<6} {statement}", flush=True)

    (out / "_timing.json").write_text(json.dumps(timing, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    await engine.dispose()
    print(f"\n用例 {len(selected)} 个，失败 {failures} 个，输出目录 {out}")
    return 1 if failures else 0


def main() -> int:
    parser = argparse.ArgumentParser(description="只读接口等价性快照")
    parser.add_argument("--cases", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--user", default="")
    parser.add_argument("--only", nargs="*", default=[])
    parser.add_argument("--page-size", type=int, default=0)
    parser.add_argument("--timeout-ms", default="300000")
    parser.add_argument("--verbose", action="store_true")
    return asyncio.run(run(parser.parse_args()))


if __name__ == "__main__":
    sys.exit(main())
